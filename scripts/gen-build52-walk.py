#!/usr/bin/env python3
"""The Build-52 founder walk sheet — what changed since Build 51, in the order you'd tap it.

One row per thing to look at, phrased as what YOU do and what you should see. Yellow columns are
yours: Pass/Fail and a comment. Nothing in here is a code instruction — it is a walk.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

OUT = 'docs/FCC-core-55-70/build52-founder-walk-2026-08-11.xlsx'
HEAD = PatternFill('solid', fgColor='085041')
YOURS = PatternFill('solid', fgColor='FFF7D6')
BAND = PatternFill('solid', fgColor='DFF2E9')
HF = Font(color='FFFFFF', bold=True, size=11)

WALK = [
    ('Your findings — naming', 'Open Net worth and read the account rows',
     'Each account reads "<institution> -<last four>" — e.g. "Vanguard -5738" — never the broker wording like "Vanguard Kamala Kavadia Brokerage". A tax-advantaged account also shows its wrapper: "Vanguard Roth IRA -5738".'),
    ('Your findings — naming', 'Find an account where the broker never shared digits',
     'It reads institution + the account word: "Vanguard Brokerage", "Chase Checking", "Ally Savings". Two that still read alike get a · 1 and a · 2.'),
    ('Your findings — banner', 'If two accounts at one firm are stale, open the missing-data banner',
     'Two DIFFERENT lines now — "Vanguard -5738 last updated Aug 3" and "Vanguard -1129 last updated Aug 3". You saw the same line twice before.'),
    ('Your findings — sync', 'Tap "Sync now" in that banner',
     'It SYNCS in place — the button shows "Syncing…" and then tells you what happened. It used to just open the account page, which cannot sync.'),
    ('Your findings — sync', 'If a connection is actually broken, look at the banner',
     '"We couldn\'t reach Vanguard -5738", the honest meanwhile, and a RECONNECT button. Before this build, a dead connection and a healthy idle one looked identical.'),
    ('Your findings — bars', 'On What you own, tap "By institution"',
     'The BAR regroups by institution too — it used to stay by category while the list below changed.'),
    ('Your findings — bars', 'Look at the What you owe bar with fresh eyes (or a greyscale screenshot)',
     'Debt segments now run dark-to-light in ONE colour, so they are told apart by lightness, not hue. Readable with any colour blindness.'),
    ('Your findings — alignment', 'Run your eye down the right edge of EVERY number on Net worth',
     'THE ONE TO CHECK HARDEST: section totals, group totals, category totals and account amounts all end on one line. Band totals used to sit 8 points right of the row values.'),
    ('Your findings — sizes', 'Compare a band total with the amounts below it',
     'All money is one size (15pt) and totals are bold. Band totals used to be smaller than row values.'),
    ('Your findings — sub-line', 'Look under an account inside Bonds & CDs',
     'It counts what is there — "3 CDs & 1 Treasury in this account" — instead of repeating the class name. Cash says "cash in this account".'),
    ('Your findings — account page', 'Open a holding account, check the two figures at the top of a class slice',
     'The line states the arithmetic: "$224,190 of $224,690 in Vanguard -5738 — the other $500 is Cash". Two bare figures used to read as a mismatch.'),
    ('Your findings — income', 'On the whole-account page, look at the income line under each holding',
     'A holding shows income ONLY when the source named that holding. No CD should ever show "dividends". Untickered income is counted once, in the Income figure at the top.'),
    ('Your findings — activity', 'Look at the ACTIVITY list on that page',
     'Newest first. Dividends and interest are visible instead of sitting past the eight-row cut, which is why the page contradicted itself before.'),
    ('Your findings — deposits', 'Record a deposit on an account that itemises holdings',
     'The sheet says "Lands as cash in this account. When you invest it, record the purchase and it moves into that holding." Then check the classification: cash goes UP, the CDs do NOT.'),
    ('Your findings — leftover', 'Compare a connected account total against its holdings',
     'Any unexplained money reads as CASH, never as more of the biggest holding. That was the chain behind your $500.'),
    ('Retirement plan', 'Open Net worth and read the retirement card',
     'It speaks even when the odds are below 80% — "Unlikely — retire at 65 with ~$X, running short around age 84". It used to fall back to a bare "See your plan".'),
    ('Retirement plan', 'If an answer is missing (age, spending, or retirement savings)',
     'You get the labelled sample plus exactly what is missing: "a sample, not your number — add what you spend and it becomes yours". Never a bare "See your plan".'),
    ('Retirement plan', 'Tap the ⓘ on the retirement card and read it end to end',
     'Plain English: what we run (500 times), what Monte Carlo means, EVERY assumption with its number, where to change them (Plan → Retirement), and the two honest limits. No jargon.'),
    ('Retirement plan', 'Change a return or inflation in Plan → Retirement',
     'The odds re-run. Risk is now measured from what you actually hold, not derived from the return.'),
    ('Banners everywhere', 'Look at the green title bars across the app',
     'Square corners, edge to edge, on EVERY screen — not just Net worth. Hero and retirement plan wear the same bar as What you own.'),
    ('Classification', 'Check where a money-market fund sits (VMFXX, SPAXX, SWVXX…)',
     'Under Stocks / ETFs, not Cash — including when the broker sends only the ticker with no name, and including holdings synced after the migration ran.'),
    ('Classification', 'Check where a CD or Treasury sits',
     'Under Bonds & CDs. Cash is cash only.'),
    ('Device-only — yours', 'Turn text size up to the largest setting',
     'The dated banner title ("YOUR NET WORTH · AUG 11, 2026") and the bar legends stay readable. NO TEST CAN CHECK THIS — it needs your eyes.'),
    ('Device-only — yours', 'Take a greyscale screenshot of Net worth',
     'Both bars still readable: every slice named with its percent, debts separated by lightness alone.'),
    ('Device-only — yours', 'Turn on reduce-motion; swipe to the app switcher',
     'Sheets appear without animation; the Net worth screen is masked in the snapshot.'),
    ('Device-only — yours', 'On the smallest phone, open the deposit sheet',
     'The keyboard never covers the save button.'),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Build 52 walk'
heads = ['#', 'Area', 'What you do', 'What you should see', 'Pass / Fail', 'Your comment']
ws.append(heads)
for c in range(1, len(heads) + 1):
    ws.cell(1, c).fill = HEAD
    ws.cell(1, c).font = HF
ws.row_dimensions[1].height = 24
for i, (area, do, see) in enumerate(WALK, 1):
    ws.append([i, area, do, see, '', ''])
for col, w in zip('ABCDEF', [5, 24, 46, 88, 12, 40]):
    ws.column_dimensions[col].width = w
for r in range(2, ws.max_row + 1):
    for c in range(1, 7):
        ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical='top')
    for c in (5, 6):
        ws.cell(r, c).fill = YOURS
    if 'Device-only' in str(ws.cell(r, 2).value):
        ws.cell(r, 2).fill = BAND
ws.freeze_panes = 'A2'

# ── the context tab ──
ws2 = wb.create_sheet('What is in this build')
rows = [
    ['Build', 'v1.1.0, cut 2026-08-11 from commit ec8e487 on branch taxonomy-v1.0.7 (pushed). The build NUMBER is read off the finished build and stated in the chat — do not assume it from the previous one.'],
    ['Since', 'Build 51 (2026-08-11, commit 8e09374) — 13 commits, all from your device findings the same day'],
    ['The headline', 'Everything you found on build 51, fixed — plus the root causes behind them. The $500 chain in particular ran three layers deep: a deposit had nowhere to land, so the total could not be explained, so the difference was guessed onto your CDs.'],
    ['Your rules recorded', 'The account-naming rule and the full retirement-odds model are now in the master PRD (FCC-core-PRD-v1.1), not only in code — Amendments tab, Sameness contract, and a new Retirement odds assumptions tab.'],
    ['Defects found while fixing', 'A right edge genuinely bent by 8 points (and the test meant to catch it could not) · income attributed by guess onto every holding · an unsorted activity list hiding the income the page counted · the cash-only rule never applied to newly-synced data · money-market funds sent as bare tickers.'],
    ['Gates at the cut', '1,522 tests green · types clean · UI-test gate passing · 5 journey suites (64 tests) green · open-gaps ledger 9 of 9 closed.'],
    ['Known and open', 'Live price provider still unselected (your call). The Vanguard income TOTAL ($27,694) — I fixed where it is attributed and that it is visible, but whether the broker sent the right rows is a data question; the activity list is now where to check it row by row.'],
    ['Not verified by anything but your eyes', 'How the dated banner title and the bar legends wrap at the largest text size on the smallest phone; and the greyscale readability of both bars on a real screen.'],
]
ws2.append(['Item', 'Detail'])
for c in (1, 2):
    ws2.cell(1, c).fill = HEAD
    ws2.cell(1, c).font = HF
for r in rows:
    ws2.append(r)
ws2.column_dimensions['A'].width = 26
ws2.column_dimensions['B'].width = 120
for r in range(2, ws2.max_row + 1):
    for c in (1, 2):
        ws2.cell(r, c).alignment = Alignment(wrap_text=True, vertical='top')
ws2.freeze_panes = 'A2'

wb.save(OUT)
print('wrote', OUT, '—', ws.max_row - 1, 'walk rows')
