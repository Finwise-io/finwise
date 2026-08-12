#!/usr/bin/env python3
"""Records everything built on 2026-08-11 in the master monitor, and points the PRD at it.

The founder's ask: "go to master list of what's missing in app versus PRD, and update it listing
what already built." The existing OPEN/PARKED rows in that file are older items today's work did not
touch — so this adds the day's build as its own tab rather than silently re-labelling unrelated rows,
and leaves every existing status exactly as it was.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

MASTER = 'docs/FCC-core-55-70/build-analysis-july24.xlsx'
PRD = 'docs/FCC-core-55-70/FCC-core-PRD-v1.1-2026-07-01.xlsx'
HEAD = PatternFill('solid', fgColor='085041')
DONE = PatternFill('solid', fgColor='DFF2E9')
OPEN_ = PatternFill('solid', fgColor='FAEEDA')
HF = Font(color='FFFFFF', bold=True, size=11)

# (Area, What was asked for / what was wrong, What is built, Where it came from, Status)
BUILT = [
    ('— NET WORTH · the Quiet Instrument rebuild —', '', '', '', ''),
    ('Net worth · the whole screen', 'Claude Design "Quiet Instrument" handoff + the founder\'s 11 hero corrections',
     'Flat continuous ledger; one shared right edge for every number; hero title carries only the info dot then the date; the trend line under the hero deleted; change line reads "Change in net worth +$X · Return on cash + investments +Y%", green at zero and up, amber only on a fall; since-line plain weight with the year.',
     'Handoff + NW-screen.docx, 2026-08-10', 'BUILT — in build 52'),
    ('Net worth · the six decisions', 'Founder decisions on the audit workbook',
     'Composition bar replaces the donut; What you owe gains its own bar; the cushion gains a progress bar; categories read biggest-first; the cash-flow row deleted; grouping buttons on day one.',
     'Audit workbook "Round 2", 2026-08-11', 'BUILT — in build 52'),
    ('All screens · section banners', 'Square corners, applied everywhere',
     'SectionBand corners squared — one component, so every banded screen changed at once. Hero and retirement plan wear the same full-width banner as WHAT YOU OWN.',
     'Founder, 2026-08-11', 'BUILT — in build 52'),
    ('All screens · the money column', 'One size, totals bold, genuinely right-aligned',
     'Every money figure 15pt, totals bold. AND the column was really bent: band totals sat 8 points right of row values because a row puts a gap before its arrow and a band did not. One constant now; the test MEASURES both edges instead of asserting a margin.',
     'Founder, 2026-08-11', 'BUILT — in build 52'),
    ('All screens · account naming', 'THE naming rule',
     'Taxable = institution + last four ("Vanguard -5738"); no digits = institution + account word; tax-advantaged = institution + wrapper + last four; duplicates get · 1 / · 2. Only real digits qualify. Recorded in the PRD (Amendments + Sameness contract).',
     'Founder rule, 2026-08-11', 'BUILT — after build 52'),

    ('— CONNECTIONS · honesty —', '', '', '', ''),
    ('Net worth · "Sync now"', 'The button did not sync — it opened a page that cannot sync',
     'It runs the sync in place, shows "Syncing…", and reports the real count. Saying "Updated" when nothing refreshed is gone.',
     'Founder device finding, 2026-08-11', 'BUILT — after build 52'),
    ('Net worth · dead connections', 'No way to re-log-in, and a dead link looked identical to a healthy one',
     'Per-account sync failures are recorded; the row reads "We couldn\'t reach Vanguard -5738" with a RECONNECT button. An account marked connected with NO connection at the provider is recorded too, instead of ageing silently.',
     'Founder device finding, 2026-08-11', 'BUILT — after build 52'),
    ('Everywhere · two dead routes', 'Buttons that led to screens which do not exist',
     'First-day "Connect a brokerage" and the fix-it sheet\'s "Update the value" both fixed; a new guard walks every pushed route and fails on any that names no screen.',
     'Found during the rebuild audit', 'BUILT — in build 52'),

    ('— CLASSIFICATION · cash means cash —', '', '', '', ''),
    ('Assets · the cash-only rule', 'Applied in the migration but NOT where data arrives',
     'The name/ticker rule is applied at classification, so a holding synced after the migration cannot land back in Cash. Known money-market tickers (VMFXX, SPAXX, SWVXX…) are recognised when a broker sends only a symbol.',
     'Founder correction, 2026-08-11', 'BUILT — after build 52'),
    ('Assets · the unexplained leftover', 'Counted as MORE of the biggest holding',
     'A positive leftover is CASH — it is money the broker\'s total includes but its holdings do not itemise. A negative leftover stays a pricing artifact. This was the chain that turned a $500 deposit into "more CDs".',
     'Founder question, 2026-08-11', 'BUILT — after build 52'),
    ('Transactions · a deposit', 'Had nowhere to land, so it only inflated the total',
     'A deposit into an account that itemises holdings opens the cash sleeve and lands there; the total rises by the same amount. The sheet says so: "Lands as cash in this account."',
     'Founder finding, 2026-08-11', 'BUILT — after build 52'),

    ('— THE RETIREMENT PLAN —', '', '', '', ''),
    ('Net worth + Plan · the verdict', 'Went silent below 80% odds — a bare "See your plan"',
     'It speaks for every computable plan and tells the truth about the ending: "to spare" only when there is something spare, otherwise "running short around age 84". The word now carries its number, in ONE wording used on Home, Net worth and Plan: "<word> — xx% odds of lasting to age abc".',
     'Founder finding + rule, 2026-08-11', 'BUILT — after build 52'),
    ('Plan · the three questions', 'A line naming one missing answer landed on a hub, then on a 7-item checklist',
     'NEW /plan-answers: asks what the odds are missing, prefilled with what we hold, ?focus= opens at one, shows the odds on the same screen, returns you where you came from. Home, Net worth and Plan all open it. The app had NO screen capturing age at all before this.',
     'Founder finding, 2026-08-11', 'BUILT — after build 52'),
    ('Plan · what counts toward retirement', 'Cash and all property were excluded by default',
     'Everything counts except the home you live in and the car you drive; a rental and a second car count. Changeable per account on the answers screen, and the person\'s choice always beats the default.',
     'Founder rule, 2026-08-11', 'BUILT — after build 52'),
    ('Plan · the odds model', '"Where do the assumptions come from?"',
     'Volatility is MEASURED per asset class from the same long-run series as the returns (was return × 1.7). The info dot explains the whole method in plain English with every number and where to change it, and it SCROLLS — the text was previously unreadable past the fold. The full model is recorded in the PRD ("Retirement odds — assumptions", 45 rows).',
     'Founder questions, 2026-08-11', 'BUILT — after build 52'),

    ('— ACCOUNT DETAIL + PERFORMANCE —', '', '', '', ''),
    ('Account detail · income', 'Untickered income was pinned onto every holding',
     'A holding claims income only when the source names it — no CD shows "dividends" any more. Untickered income is still counted at the account level. The activity list is newest-first, so income is no longer hidden past the cut.',
     'Founder finding, 2026-08-11', 'BUILT — after build 52'),
    ('Account detail · whose fact is a row', 'Broker records and your own looked identical',
     'Your rows wear "you entered this" and a pencil; broker rows read "from Vanguard" and are read-only. A correction reverses the old amount and records the new one through the one engine. The door names the date the list stops at.',
     'Performance handoff, 2026-08-11', 'BUILT — after build 52'),
    ('Performance · the walks', 'Per-account walks skipped accounts with no itemised holdings',
     'Every account with value gets its walk, so the parts sum to the all-accounts walk — which is titled "· ALL ACCOUNTS" as the handoff names it.',
     'Performance handoff, 2026-08-11', 'BUILT — after build 52'),
    ('Performance · still open', 'Two items deliberately not built',
     'Category rows tapping through to a category-detail screen (a new screen). The handoff\'s "/yr" stated-yield treatment for bonds and cash (a CALCULATION change — the founder\'s instruction was to take the UX, not the calculations).',
     'Performance handoff, 2026-08-11', 'OPEN — founder\'s call'),
]


def build_tab():
    wb = openpyxl.load_workbook(MASTER)
    title = 'Built — Aug 11'
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title, 1)
    heads = ['Area', 'What was asked for / what was wrong', 'What is built', 'Where it came from', 'Status']
    ws.append(heads)
    for c in range(1, 6):
        ws.cell(1, c).fill = HEAD
        ws.cell(1, c).font = HF
    ws.row_dimensions[1].height = 22
    for row in BUILT:
        ws.append(list(row))
        r = ws.max_row
        for c in range(1, 6):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical='top')
        if row[0].startswith('—'):
            for c in range(1, 6):
                ws.cell(r, c).fill = DONE
                ws.cell(r, c).font = Font(bold=True)
        elif str(row[4]).startswith('OPEN'):
            for c in range(1, 6):
                ws.cell(r, c).fill = OPEN_
    for col, w in zip('ABCDE', [34, 54, 94, 32, 26]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'

    # a pointer row on the summary tab, so the day is visible from the front of the file
    sm = wb['SUMMARY — read me first']
    sm.append(['2026-08-11', 'Built — Aug 11 (this file)',
               'The Net worth Quiet-Instrument rebuild, the founder\'s device findings from builds 51 and 52, and the Performance/account-detail handoff.',
               'See the "Built — Aug 11" tab: 18 items with what was wrong, what is built and where it came from.',
               'BUILT (build 52 + after)',
               'Two Performance items remain open by the founder\'s own instruction — a category-detail screen, and the "/yr" stated-yield treatment (a calculation change).'])
    for c in range(1, 7):
        sm.cell(sm.max_row, c).alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(MASTER)
    print('master monitor: "Built — Aug 11" tab written,', len(BUILT), 'rows')


def prd_pointer():
    wb = openpyxl.load_workbook(PRD)
    am = wb['Amendments v1.1→v1.2']
    am.append(['2026-08-11', 'Build status — what is built vs what the documents ask for',
               'The 2026-08-11 build is recorded in the master monitor (docs/FCC-core-55-70/build-analysis-july24.xlsx, tab "Built — Aug 11"): 18 items covering the Net worth Quiet-Instrument rebuild, the connection-honesty fixes, the classification chain (cash means cash; the leftover is cash; a deposit lands in cash), the retirement plan (the verdict speaks below 80%, the three-questions screen, what counts toward retirement, measured volatility, the plain-English method), and the account-detail / Performance handoff. Each row carries what was wrong, what is built, and where the instruction came from. TWO items remain open by the founder\'s own instruction: a category-detail screen, and the handoff\'s "/yr" stated-yield treatment for bonds and cash, which is a calculation change rather than UX.',
               'Founder 2026-08-11: "Update master PRD document. Go to master list of what\'s missing in app versus PRD, and update it listing what already build."'])
    for c in range(1, 5):
        am.cell(am.max_row, c).alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(PRD)
    print('PRD: Amendments row added — now', am.max_row - 1, 'rows')


build_tab()
prd_pointer()
