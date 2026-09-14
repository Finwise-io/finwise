"""B2 — the formula in the code and the formula in the file must say the same thing.

The founder's rule, made mechanical: *if the inputs, the variables, the assumptions and the formula
are the same, the output is the same.* So the arithmetic printed above a function in the code and
the arithmetic on its row in the master file are two statements of one rule. They are compared as
they sit ON DISK — not as two things the same generator just produced — because a guard that reads
one source and checks it against itself proves nothing.

    python3 "Calculation Module code/scripts/check-formula-agreement.py"

Fails the build on any disagreement. Run from the repository root.
"""
import os, re, sys
import openpyxl

MASTER = 'MoneyKeel/documents/other/moneykeel_calc_master_calculation-master_v01.xlsx'
import importlib.util as _il, os as _os
_spec=_il.spec_from_file_location('mkfn', _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), 'moneykeel_functions.py')); _mkfn=_il.module_from_spec(_spec); _spec.loader.exec_module(_mkfn)
ROOTS = _mkfn.CALC_ROOTS   # 2026-09-12: the ONE roots list — the sweep, the header writer and this check read the same files
HEAD   = re.compile(r'^/\* ─ (MK-[A-Z]{2}-[FO]-\d{3}) ─')


def from_code():
    """{function id: the arithmetic printed above it}, read out of the source files."""
    out = {}
    for root in ROOTS:
        for dp, _, fs in os.walk(root):
            if '__tests__' in dp:
                continue
            for f in fs:
                if not f.endswith(('.ts', '.tsx')) or '.test.' in f:
                    continue
                lines = open(os.path.join(dp, f)).read().split('\n')
                i = 0
                while i < len(lines):
                    m = HEAD.match(lines[i])
                    if not m:
                        i += 1; continue
                    body, j = [], i + 1
                    while j < len(lines) and not lines[j].strip().startswith('*/'):
                        body.append(lines[j]); j += 1
                    take, got = False, []
                    for b in body:
                        t = re.sub(r'^\s*\*\s?', '', b)
                        if t.startswith('Formula .....'):
                            take = True; continue
                        if take:
                            if not t.strip() or re.match(r'^[A-Z][a-z]+ \.+', t):
                                take = False; continue
                            got.append(t.rstrip())
                    out[m.group(1)] = '\n'.join(x[4:] if x.startswith('    ') else x
                                                for x in got).strip()
                    i = j
    return out


def from_file():
    """{function id: the arithmetic on its row}, read out of the master file."""
    wb = openpyxl.load_workbook(MASTER, data_only=True)
    ws = wb['functions_from_code']
    head = [str(c.value) for c in ws[1]]
    fi = head.index('Formula (mathematical)')
    out = {r[0]: str(r[fi] or '').strip() for r in ws.iter_rows(min_row=3, values_only=True) if r[1]}
    wb.close()
    return out


def norm(t):
    """Whitespace is not a disagreement; anything else is."""
    return re.sub(r'\s+', ' ', str(t)).strip()


if __name__ == '__main__':
    code, filed = from_code(), from_file()
    both = sorted(set(code) & set(filed))
    bad = [i for i in both if norm(code[i]) != norm(filed[i])]
    only_code = sorted(set(code) - set(filed))
    only_file = sorted(i for i in set(filed) - set(code) if filed[i])

    print('formula agreement — %d functions carry the rule in both places' % len(both))
    if only_code:
        print('  %d have a header but no row in the file' % len(only_code))
    if only_file:
        print('  %d have a row in the file but no header in the code' % len(only_file))
    if not bad:
        print('  every one of the %d agrees.' % len(both))
        sys.exit(0)
    print('\nBUILD STOPPED — %d disagree. There is no winner: read both, find which side is wrong, '
          'and correct both.\n' % len(bad))
    for i in bad[:12]:
        print('  %s' % i)
        print('    the code says : %s' % norm(code[i])[:150])
        print('    the file says : %s' % norm(filed[i])[:150])
    sys.exit(1)
