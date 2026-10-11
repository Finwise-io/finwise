#!/usr/bin/env python3
"""AUDIT ONLY (2026-10-11) — audit_inputs_and_use: writes the founder's workbook from the rows that
audit_inputs_and_use.py produced. Changes no app code.

    python3 audit_inputs_and_use_workbook.py ROWS.json OUT.xlsx
"""
import json, os, sys, collections
from openpyxl import Workbook
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from xlsx_readable import readable

R = json.load(open(sys.argv[1])); OUT = sys.argv[2]
ORDER = {'Net worth': 0, 'Performance': 1, 'Cash flow': 2}
def mkey(r, idk): return (ORDER.get(r.get('Module'), 9), r.get('Module', ''), r[idk])

DEFAULT_PROBLEM = {
    'TESTS ONLY': 'Only tests run it; no screen does. Its tests pass on a number nobody sees.',
    'UNUSED': 'Nothing in the code calls it.',
    'USED ONLY BY OTHER UNUSED FUNCTIONS': 'Its only callers are themselves unused, so it never runs in the app.',
}
wb = Workbook()

# ── Inputs ──
ws = wb.active; ws.title = 'Inputs'
cols = ['Function id', 'Function name', 'Input', 'Origin', 'Evidence', 'Problem', 'Module', 'Function used?']
ws.append(cols)
inputs = sorted(R['inputs'], key=lambda r: (ORDER.get(r['Module'], 9), r['Module'], r['Function id']))
for r in inputs:
    prob = r['Problem']
    if r['Used?'] != 'USED' and r['Origin'] in ('MISSING', 'DECIDED?') and prob:
        prob = f"(the function itself is {r['Used?'].lower()}, so this does not reach a screen) " + prob
    ws.append([r['Function id'], r['Function name'], r['Input'], r['Origin'], r['Evidence'], prob, r['Module'], r['Used?']])
readable(ws)

# ── Use ──
ws = wb.create_sheet('Use')
ws.append(['Id', 'Name', 'Kind', 'Used?', 'Where', 'Problem', 'Module', 'Code'])
use = sorted(R['use'], key=lambda r: mkey(r, 'Id')) + sorted(R['metrics'], key=lambda r: mkey(r, 'Id'))
for r in use:
    prob = r['Problem'] or DEFAULT_PROBLEM.get(r['Used?'], '')
    ws.append([r['Id'], r['Name'], r['Kind'], r['Used?'], r['Where'], prob, r['Module'], r.get('_file', '')])
readable(ws)

# ── Inputs register ──
ws = wb.create_sheet('Inputs register')
ws.append(['Input', 'Name', 'Source (register)', 'Fields', 'Origin found', 'Evidence', 'Read by functions', 'Readers (first)'])
for r in R['register']:
    ws.append([r['Input'], r['Name'], r['Source (register)'], r['Fields'], r['Origin found'], r['Evidence'], r['Read by functions'], r['Readers (first)']])
readable(ws)

# ── Entered, read by no function ──
ws = wb.create_sheet('Entered not read')
ws.append(['Field', 'Entered at', 'Read by'])
for r in R['entered_unused']:
    ws.append([r['Field'], r['Entered at'], r['Read by']])
readable(ws)

# ── Coverage ──
ws = wb.create_sheet('Coverage')
c = R['counts']
uc = collections.Counter(r['Used?'] for r in R['use']); mc = collections.Counter(r['Used?'] for r in R['metrics'])
oc = collections.Counter(r['Origin'] for r in R['inputs'])
ws.append(['What', 'Count'])
for k, v in [('function and orchestrator ids issued in the register', c['functions']), ('of them found in the code by header', c['in_code']),
             ('metrics in force (register, less retired)', c['metrics']), ('input rows (function × input)', c['inputs_rows']),
             ('inputs on the inputs register', c['register_inputs']), ('routes on disk', c['routes']), ('routes the live app can reach', c['live_routes']),
             ('routes with no way in', ', '.join(R['dead_routes']))]:
    ws.append([k, v])
ws.append(['', ''])
for k, v in uc.most_common(): ws.append([f'functions/orchestrators — {k}', v])
for k, v in mc.most_common(): ws.append([f'metrics — {k}', v])
for k, v in oc.most_common(): ws.append([f'input rows — {k}', v])
readable(ws, filt=False)
wb.save(OUT)
print(f'wrote {OUT}: Inputs {len(inputs)} rows, Use {len(use)} rows, register {len(R["register"])}, entered-not-read {len(R["entered_unused"])}')
