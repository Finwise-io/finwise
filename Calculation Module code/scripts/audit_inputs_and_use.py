#!/usr/bin/env python3
"""AUDIT ONLY (2026-10-11) — audit_inputs_and_use: for the founder's two questions.

  (a) every input a function / orchestrator reads is entered by a person, set by his decision, a
      published figure, delivered by a feed, or made by another function — nothing is MISSING;
  (b) every function, orchestrator and metric is used by the live app.

Reads the graph written by audit_inputs_and_use_graph.cjs (the TypeScript compiler's own reading of
the code: who calls whom, what each function reads, where every field is written) plus the
registers. Writes JSON rows for the workbook to the path given. Changes no app code.

    node audit_inputs_and_use_graph.cjs /tmp/graph.json
    python3 audit_inputs_and_use.py /tmp/graph.json /tmp/rows.json

Every scan states how many records it saw and asserts a floor (CLAUDE.md, A GUARD MUST PROVE IT LOOKED).
"""
import json, os, re, sys, collections, importlib.util

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
G = json.load(open(sys.argv[1]))
OUT = sys.argv[2]

def load(name): return json.load(open(os.path.join(HERE, name)))
REG = load('moneykeel_function_register.json')
MREG = load('moneykeel_metric_register.json')
FIGS = load('moneykeel_figures.json')['figures']
CENSUS = load('moneykeel_figure_census.json')
PUB = {r['symbol']: r for r in load('moneykeel_published_sources.json')}
DECS = load('moneykeel_decisions.json')['rows']
_spec = importlib.util.spec_from_file_location('mki', os.path.join(HERE, 'moneykeel_inputs.py'))
MKI = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(MKI)
_spec = importlib.util.spec_from_file_location('mkcf', os.path.join(HERE, 'moneykeel_code_formula.py'))
CF = importlib.util.module_from_spec(_spec)
try:
    _spec.loader.exec_module(CF); ASSUME = getattr(CF, 'ASSUMPTIONS_EXTRA', {})
except Exception:
    ASSUME = {}

U = G['units']; EDGES = G['edges']; FACTS = G['facts']
WRITES = G['writes']; READS = G['reads']; MENTIONS = G['mentions']; CALLS = G['calls']; STORECALLS = G['storeCalls']
print(f"graph: {G['counts']}")
assert G['counts']['files'] >= 600 and G['counts']['registered'] >= 700, 'the graph saw too little'

# ── the inputs register (master file tab "inputs" + moneykeel_inputs.NEW) ───────────────────────
import openpyxl
wb = openpyxl.load_workbook(os.path.join(ROOT, 'MoneyKeel/documents/other/moneykeel_calc_master_calculation-master_v01.xlsx'), read_only=True)
INPUTS = {}
for i, r in enumerate(wb['inputs'].iter_rows(values_only=True)):
    if i == 0 or not r[0] or not str(r[0]).startswith('MK-XS-I-'): continue
    INPUTS[r[0]] = {'name': r[1], 'source': r[2] or '', 'fields': r[3] or '', 'active': r[7]}
for k, (nm, src, flds) in MKI.NEW.items():
    INPUTS.setdefault(k, {'name': nm, 'source': src, 'fields': ' · '.join(flds), 'active': 'ACTIVE'})
INPUT_FIELDS = {k: v for k, v in list(MKI.FIELDS.items()) + [(k, v[2]) for k, v in MKI.NEW.items()]}
print(f"inputs register: {len(INPUTS)} inputs"); assert len(INPUTS) >= 60

FIELD_TO_INPUT = {}
for iid, flds in INPUT_FIELDS.items():
    for f in flds:
        clean = f.replace('\\b', '').replace('\\', '')
        if re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*', clean): FIELD_TO_INPUT.setdefault(clean, iid)
        m = re.fullmatch(r'[a-z]+\\?\.([A-Za-z_]+)', f)
        if m: FIELD_TO_INPUT.setdefault(m.group(1), iid)

# ── registered ids ─────────────────────────────────────────────────────────────────────────────
ISSUED = REG['issued']; NAME_OF = {v: k for k, v in ISSUED.items()}
REGU = {u['id']: u for u in U if u.get('id')}
print(f"function register: {len(NAME_OF)} ids issued; {len(REGU)} found in code by header")
assert len(REGU) >= 720

def module_of(fid):
    fam = fid[3:5]
    return {'NW': 'Net worth', 'PF': 'Performance', 'CF': 'Cash flow', 'PL': 'Plan', 'HM': 'Home', 'MN': 'Menu', 'ST': 'Storage',
            'XS': 'Cross-screen'}.get(fam, fam)
SCOPE_ORDER = {'Net worth': 0, 'Performance': 1, 'Cash flow': 2}

# ── header parsing ─────────────────────────────────────────────────────────────────────────────
LABEL = re.compile(r'^ \* ([A-Z][A-Za-z ]+?) \.{2,}\s?(.*)$')
def header_fields(h):
    out = collections.OrderedDict(); cur = None
    for ln in (h or '').split('\n')[2:]:
        m = LABEL.match(ln)
        if m: cur = m.group(1).strip(); out[cur] = [m.group(2).strip()] if m.group(2).strip() else []; continue
        if cur and ln.startswith(' *'): out[cur].append(ln[2:].strip())
    return {k: '\n'.join(x for x in v if x) for k, v in out.items()}
HF = {fid: header_fields(u['header']) for fid, u in REGU.items()}

# ── file classes ───────────────────────────────────────────────────────────────────────────────
def is_test(f): return bool(re.search(r'__tests__|\.test\.tsx?$|\.probe\.|scratch_tests|jest\.setup|/bench/', f))
def is_demo(f, unit=None):
    if re.search(r'src/testing/|persona|fixture|sample|demo', f, re.I): return True
    if unit is not None:
        nm = U[unit]['name'] or ''
        if re.search(r'EXAMPLE|WORKED|SAMPLE|DEMO|PERSONA|FIXTURE|example', nm): return True
        uid = U[unit].get('id')
        if uid and 'Example only' in HF.get(uid, {}).get('Answers', ''): return True
    return False
def fclass(f):
    if is_test(f): return 'test'
    if f.startswith('app/') or f.startswith('src/screens/') or f.startswith('src/components/') or f.startswith('src/onboarding/'): return 'screen'
    if re.match(r'src/services/(sync/|marketData|economicData|receipt)|src/domain/import/', f): return 'feed'
    if f.startswith('src/store/'): return 'store'
    return 'calc'

# ── graph ──────────────────────────────────────────────────────────────────────────────────────
OUTE = collections.defaultdict(set); INE = collections.defaultdict(set)
for a, b in EDGES:
    OUTE[a].add(b); INE[b].add(a)
BY_FILE = collections.defaultdict(list)
for u in U: BY_FILE[u['file']].append(u['uid'])
MODULE_UNIT = {u['file']: u['uid'] for u in U if u['name'] == '<module>'}

def is_example_unit(uid):
    u = U[uid]
    if u.get('id') and 'Example only' in HF.get(u['id'], {}).get('Answers', ''): return True
    return bool(re.search(r'^(EXAMPLE|WORKED)|_EXAMPLE\b|EXAMPLE_', u['name'] or ''))
def bfs(starts):
    seen = set(starts); stack = list(starts); files = set()
    while stack:
        x = stack.pop(); f = U[x]['file']
        if f not in files:
            files.add(f); m = MODULE_UNIT.get(f)
            if m is not None and m not in seen: seen.add(m); stack.append(m)
        for y in OUTE[x]:
            if y not in seen and not U[y]['test'] and not is_example_unit(y): seen.add(y); stack.append(y)
    return seen, files

# ── live routes ────────────────────────────────────────────────────────────────────────────────
ROUTES = {}
for f in sorted(BY_FILE):
    m = re.fullmatch(r'app/(\(tabs\)/)?([a-z0-9-]+|index)\.tsx', f)
    if m and m.group(2) != '_layout': ROUTES[('/(tabs)/' if m.group(1) else '/') + m.group(2)] = f
print(f"routes on disk: {len(ROUTES)}"); assert len(ROUTES) >= 40
tabs_src = open(os.path.join(ROOT, 'src/constants/tabs.ts')).read()
TAB_BAR = re.findall(r'^\s{2}(\w+):\s*\{\s*title:\s*\'([^\']+)\'', tabs_src, re.M)
TAB_TITLE = dict(TAB_BAR)
assert len(TAB_BAR) == 5, TAB_BAR
NO_DOOR = {'/index', '/auth', '/first-run', '/onboarding', '/connect-done'}
stack_leaves = {r[1:] for r in ROUTES if not r.startswith('/(tabs)/')}
def door_patterns(route):
    leaf = route.split('/')[-1]
    forms = [route] if not route.startswith('/(tabs)/') else [route] + ([f'/{leaf}'] if leaf not in stack_leaves else [])
    pats = []
    for fm in forms: pats += [f"'{fm}'", f'"{fm}"', f'`{fm}', f"'{fm}?", f'"{fm}?', f"'{fm}/"]
    return pats
def strip_comments(t):
    t = re.sub(r'/\*.*?\*/', lambda m: '\n' * m.group(0).count('\n'), t, flags=re.S)
    return re.sub(r'(^|[^:\'"`])//[^\n]*', r'\1', t)
SRC = {}
def src(f):
    if f not in SRC:
        try: SRC[f] = strip_comments(open(os.path.join(ROOT, f)).read())
        except Exception: SRC[f] = ''
    return SRC[f]
ROOT_FILES = ['app/_layout.tsx', 'app/(tabs)/_layout.tsx']
live = {f'/(tabs)/{k}' for k, _ in TAB_BAR} | (NO_DOOR & set(ROUTES))
door_of = {r: ('in the tab bar' if r.startswith('/(tabs)/') else 'reached without a tap') for r in live}
while True:
    starts = [x for f in ROOT_FILES + [ROUTES[r] for r in live] for x in BY_FILE.get(f, [])]
    reached, rfiles = bfs(starts)
    added = False
    for r in ROUTES:
        if r in live: continue
        pats = door_patterns(r)
        for f in sorted(rfiles):
            if 'routeGuard' in f or f == ROUTES[r] or is_test(f): continue
            hit = None
            for p in pats:
                for m in re.finditer(re.escape(p), src(f)):
                    t = src(f); ls = t.rfind('\n', 0, m.start()) + 1
                    before = t[ls:m.start()]
                    # a figure's own list of the screens it is printed on is not a way in
                    if re.search(r'\bscreens\s*:\s*\[', before) or re.search(r'\broutes\s*:\s*\[', before): continue
                    hit = (f, t.count('\n', 0, m.start()) + 1); break
                if hit: break
            if hit:
                live.add(r); door_of[r] = f'{hit[0]}:{hit[1]}'; added = True; break
    if not added: break
DEAD_ROUTES = sorted(set(ROUTES) - live)
print(f"live routes: {len(live)} of {len(ROUTES)}; no way in: {DEAD_ROUTES}")
REACHED, RFILES = reached, rfiles

def route_name(r):
    leaf = r.split('/')[-1]
    if r.startswith('/(tabs)/') and leaf in TAB_TITLE: return f'{TAB_TITLE[leaf]} tab ({r})'
    return r
PER_ROUTE = {}
for r in sorted(live):
    s, fs = bfs(BY_FILE.get(ROUTES[r], []))
    PER_ROUTE[r] = (s, fs)
LAYOUT_SEEN, _ = bfs([x for f in ROOT_FILES for x in BY_FILE.get(f, [])])
# what each dead route reaches (for "used only by other unused")
PER_DEAD = {r: bfs(BY_FILE.get(ROUTES[r], [])) for r in DEAD_ROUTES}

def where_uid(uid):
    rs = [route_name(r) for r in sorted(PER_ROUTE) if uid in PER_ROUTE[r][0]]
    if not rs and uid in LAYOUT_SEEN: rs = ['app start (the root layout, every screen)']
    return rs
def where_file(f):
    return [route_name(r) for r in sorted(PER_ROUTE) if f in PER_ROUTE[r][1]] or (['app start (the root layout)'] if f in bfs([x for ff in ROOT_FILES for x in BY_FILE.get(ff, [])])[1] else [])

# ── PART B — use of every function / orchestrator ──────────────────────────────────────────────
def uname(uid):
    u = U[uid]
    return f"{u['id']} {u['name']}" if u.get('id') else f"{u['name']} ({u['file']})"
use_rows = []
USE = {}
for fid in sorted(NAME_OF):
    nm = NAME_OF[fid]; kind = 'orchestrator' if '-O-' in fid else 'function'
    u = REGU.get(fid)
    if not u:
        use_rows.append({'Id': fid, 'Name': nm, 'Kind': kind, 'Module': module_of(fid), 'Used?': 'NOT IN CODE',
                         'Where': 'no header for this id anywhere in the code', 'Problem': 'The register issues this id but no function in the code carries its header — it cannot be traced.'})
        USE[fid] = 'NOT IN CODE'; continue
    uid = u['uid']; hf = HF[fid]
    example = 'Example only' in hf.get('Answers', '')
    callers = [a for a in INE[uid] if U[a]['file'] != u['file'] or U[a]['uid'] != uid]
    nontest = [a for a in callers if not U[a]['test']]
    tests = [a for a in callers if U[a]['test']]
    where = where_uid(uid)
    problem = []
    if uid in REACHED:
        cls = 'EXAMPLE ONLY' if example else 'USED'
        w = '; '.join(where)
        if example: problem.append('Header says example only, yet a live screen reaches it — check it only shows the worked example.')
    elif example:
        cls = 'EXAMPLE ONLY'; w = 'worked example; ' + (f'{len(tests)} test file(s) run it' if tests else 'nothing runs it')
    elif nontest:
        cls = 'USED ONLY BY OTHER UNUSED FUNCTIONS'
        dead = [route_name(r) for r in DEAD_ROUTES if uid in PER_DEAD[r][0]]
        names = sorted({uname(a) for a in nontest})[:6]
        w = 'called by ' + ' · '.join(names) + (f' — reached only from screen(s) with no way in: {", ".join(dead)}' if dead else '')
    elif tests:
        cls = 'TESTS ONLY'; w = f'{len({U[a]["file"] for a in tests})} test file(s): ' + ', '.join(sorted({U[a]['file'].split('/')[-1] for a in tests})[:4])
    else:
        cls = 'UNUSED'; w = 'nothing in the code calls it'
    serves = hf.get('Serves', '')
    served = re.findall(r'MK-[A-Z]{2}-M-\d{3}', serves)
    USE[fid] = cls
    use_rows.append({'Id': fid, 'Name': nm, 'Kind': kind, 'Module': module_of(fid), 'Used?': cls, 'Where': w, 'Problem': ' '.join(problem),
                     '_served': served, '_file': f"{u['file']}:{u['line']}", '_uid': uid, '_noserve': 'NO FIGURE ON RECORD' in serves})
print('use classes (functions+orchestrators):', collections.Counter(USE.values()))

# ── metrics ────────────────────────────────────────────────────────────────────────────────────
ret_m = MREG['retired']
metric_key = {}
for k, v in MREG['issued'].items():
    if '-M-' in v: metric_key.setdefault(v, []).append(k)
METRICS = sorted(set(metric_key) - set(ret_m))
print(f"metric register: {len(set(metric_key))} ids issued, {len(set(metric_key) & set(ret_m))} retired, {len(METRICS)} in force; figure register {len(FIGS)}")
assert len(METRICS) >= 200
census_seen = collections.defaultdict(set)
for route, states in CENSUS['declared'].items():
    for st, ids in states.items():
        for i in ids: census_seen[i].add(f'{route} ({st})')
# every place an id is written as a string in the code (not comments, not tests)
ID_SITES = collections.defaultdict(list)
for f in BY_FILE:
    if is_test(f): continue
    t = src(f)
    for m in re.finditer(r'[\'"`](MK-[A-Z]{2}-M-\d{3})[\'"`]', t):
        before = t[max(0, m.start() - 40):m.start()]
        # a PRINT: the id handed to fig()/figSaid()/figPer()/figWords()/figSpoken(), or written on a
        # screen or component. A door's own declaration of the figure is not a print.
        if fclass(f) == 'screen' or re.search(r'\bfig\w*\(\s*$', before):
            ID_SITES[m.group(1)].append((f, t.count('\n', 0, m.start()) + 1))
SPEC_FILES = {'MoneyKeel/functions/figures.ts', 'MoneyKeel/functions/_figure-spec.ts'}
SERVED_BY = collections.defaultdict(list)
for r in use_rows:
    for m in r.get('_served', []): SERVED_BY[m].append(r['Id'])
metric_rows = []; MUSED = {}
for mid in METRICS:
    fig = FIGS.get(mid)
    keys = metric_key[mid]
    label = fig['label'] if fig else keys[0].split('|')[-2] if keys[0].split('|')[2] != 'declared' else keys[0].split('|')[-1]
    if not fig and keys[0].split('|')[2] == 'declared': label = keys[0].split('|', 3)[3]
    sites = [(f, l) for f, l in ID_SITES.get(mid, []) if f not in SPEC_FILES and not is_demo(f)]
    # a declaration inside a door counts only as "the door declares it"; a print is fig*/figure call
    live_sites = [(f, l) for f, l in sites if f in RFILES]
    screens = sorted({w for f, l in live_sites for w in where_file(f)})
    seen = sorted(census_seen.get(mid, []))
    problem = ''
    if live_sites:
        cls = 'PRINTED'
        w = 'printed through its identifier on: ' + '; '.join(screens[:8]) + f' — e.g. {live_sites[0][0]}:{live_sites[0][1]}'
        if seen: w += f' · seen rendered: {", ".join(seen[:4])}'
        elif fig: problem = 'Not seen in any of the four sample-household renders of the figure census — printed only in some data states, or the census does not open that screen.'
    elif sites:
        cls = 'NOT PRINTED ON A LIVE SCREEN'
        w = 'its identifier appears only in: ' + '; '.join(f'{f}:{l}' for f, l in sites[:4])
        problem = 'The only code that prints it is not reachable from any live screen.'
    else:
        # hand-registered by its words: look for the expression the key names on the screen named
        parts = keys[0].split('|')
        route = parts[2] if parts[2].startswith('/') else None
        expr = parts[-1]
        hits = []
        rt = None
        if route:
            rt = route if route in ROUTES else ('/(tabs)' + route if '/(tabs)' + route in ROUTES else ('/(tabs)/analytics' if route == '/networth' else None))
        pool = sorted(PER_ROUTE[rt][1]) if rt in PER_ROUTE else ([] if rt else sorted(RFILES))
        generic = len(expr) < 10 and '.' not in expr
        for f in pool:
            if fclass(f) != 'screen': continue
            parts = [x.strip() for x in re.split(r'\s+/\s+', expr) if x.strip()]
            if expr and len(expr) > 3 and all(re.search(r'(?<![\w.])' + re.escape(x) + r'(?![\w])', src(f)) for x in parts): hits.append(f)
        weak = bool(hits and generic and not rt)
        if rt and rt not in PER_ROUTE and not fig:
            cls = 'NOT PRINTED ON A LIVE SCREEN'; w = f'its screen {route} has no way in (or no route of that name)'
            problem = 'The screen this figure is on is not reachable in the live app.'
        elif fig:
            cls = 'NOT PRINTED UNDER ITS OWN ID'; w = 'declared in the figure register (screens: ' + ', '.join(fig.get('screens', [])) + ') but no screen code prints this identifier'
            problem = 'No screen prints this identifier. Either the screen no longer shows it, or it shows the number under another id or with none — the trail from screen to formula is broken either way.'
        elif hits:
            cls = 'PRINTED (matched by the words, not by identifier)'
            w = f"the expression `{expr}` is in {', '.join(hits[:3])}" + (f' — on {route_name(rt)}' if rt else ' — ' + '; '.join(sorted({x for f in hits for x in where_file(f)})[:5]))
            if rt and rt not in live: cls = 'NOT PRINTED ON A LIVE SCREEN'
            if weak: cls = 'UNSETTLED'; problem = f'Only a common word (`{expr}`) ties it to a screen — not settled; it carries no identifier.'
            problem = problem or 'Printed without its identifier, so nothing can check it (no fig() call).'
        else:
            cls = 'NOT PRINTED ANYWHERE'
            w = f'registered by its words ({keys[0][:120]}); the expression `{expr}` is on no live screen'
            problem = 'No live screen prints it.'
    if not SERVED_BY.get(mid) and cls.startswith('PRINTED'):
        problem = (problem + ' ' if problem else '') + 'No function header names it under "Serves", so the trail from screen to formula is broken.'
    MUSED[mid] = cls
    metric_rows.append({'Id': mid, 'Name': label, 'Kind': 'metric', 'Module': module_of(mid), 'Used?': cls, 'Where': w, 'Problem': problem})
print('metric classes:', collections.Counter(MUSED.values()))

# functions that serve figures none of which are printed
for r in use_rows:
    sv = r.get('_served') or []
    if sv:
        printed = [m for m in sv if MUSED.get(m, '').startswith('PRINTED')]
        if not printed:
            r['Problem'] = (r['Problem'] + ' ' if r['Problem'] else '') + f"Its header says it serves {', '.join(sv[:5])}, and none of those is printed on a live screen."
    elif r.get('_noserve') and r['Used?'] == 'USED' and '-O-' in r['Id']:
        r['Problem'] = (r['Problem'] + ' ' if r['Problem'] else '') + 'Header names no figure it serves.'

# ── PART A — inputs ────────────────────────────────────────────────────────────────────────────
CITE_DEC = re.compile(r'DEC-\d{4}|\bG\d{1,2}\b|\bQ-?\d{1,3}\b|\bA\d{2}\b|\bB\d{1,2}\b|(?i:founder|ruling|ruled|his words|approved|owner decision|decision)')
CITE_PUB = re.compile(r'CFPB|FICO|S&P 500|Treasury|IRS|SSA|Social Security Act|Rev\. ?Proc|IRC|§|Treas|Publication \d|Notice 20|SECURE|BLS|Bureau of Labor|CPI|statut|Pub\. ?\d|Medicare|FINRA|SEC rule|Fed(eral Reserve)?')
def decs_naming(term, file=None):
    out = []
    for d in DECS:
        if d.get('status') == 'SUPERSEDED': continue
        blob = ' '.join(str(d.get(k) or '') for k in ('plain', 'his_words', 'means', 'raw', 'search', 'governs', 'changed'))
        if re.search(r'\b' + re.escape(term) + r'\b', blob) and (file is None or file in blob): out.append(d['id'])
    return out
def file_lines(f):
    try: return open(os.path.join(ROOT, f)).read().split('\n')
    except Exception: return []

def writes_of(field, exclude_uid=None):
    ws = [w for w in WRITES.get(field, []) if not is_test(w['file'])]
    real = [w for w in ws if not is_demo(w['file'], w['unit'])]
    demo = [w for w in ws if is_demo(w['file'], w['unit'])]
    return real, demo
def classify_field(field, reader_uid=None):
    """origin of a stored field, from where the code writes it"""
    real, demo = writes_of(field)
    if not real:
        return ('MISSING', ('written only in example / sample data: ' + ', '.join(f"{w['file']}:{w['line']}" for w in demo[:3])) if demo
                else 'no code anywhere writes this field (object key, assignment or string key) outside tests')
    SAVE = re.compile(r'^(set|add|update|save|record|confirm|seed|commit|persist|write|patch|put|fund|allocate|import|on[A-Z]|setState|push|ingest)', re.I)
    def entry_like(w):
        if w['how'] == 'assignment': return 2
        if w.get('jsx') in ('k', 'field', 'name', 'answerKey', 'fieldKey', 'valueKey'): return 2
        if w.get('callee') and (SAVE.match(w['callee']) or w['callee'] == 'setAnswer'): return 2
        if w.get('other') in ('VariableDeclaration', 'ReturnStatement', 'ArrowFunction'): return 1
        return 0
    screen = sorted([w for w in real if fclass(w['file']) == 'screen' and entry_like(w)], key=lambda w: -entry_like(w))
    feed = [w for w in real if fclass(w['file']) == 'feed']
    store = [w for w in real if fclass(w['file']) == 'store']
    calc = [w for w in real if fclass(w['file']) == 'calc']
    own = [w for w in real if w['unit'] == reader_uid]
    live_screen = [w for w in screen if w['file'] in RFILES]
    def ev(ws, n=2): return ', '.join(f"{w['file']}:{w['line']} ({w['how']}" + (f" in {w['callee']}()" if w.get('callee') else '') + (f" {w['jsx']}=" if w.get('jsx') else '') + ')' for w in ws[:n])
    if live_screen:
        o = 'ENTERED' + (' or FEED' if feed else '')
        return (o, 'set on screen: ' + ev(live_screen) + (('; feed: ' + ev(feed, 1)) if feed else ''))
    if feed:
        return ('FEED', 'written by the feed / download reader: ' + ev(feed) + (('; also screen code not reachable: ' + ev(screen, 1)) if screen else ''))
    if screen:
        return ('MISSING', 'written only on screen code no live route reaches: ' + ev(screen))
    if calc or store:
        ws = calc + store
        if own and len(own) == len(ws):
            return ('DERIVED', 'made inside this function itself: ' + ev(own, 1))
        wu = [w for w in ws if w['unit'] is not None]
        ids = sorted({U[w['unit']]['id'] for w in wu if U[w['unit']].get('id')})
        live_w = [w for w in ws if w['unit'] in REACHED]
        tag = ('DERIVED', 'written by ' + (', '.join(ids[:3]) + ' — ' if ids else '') + ev(ws))
        if not live_w:
            return ('MISSING', 'written only by code no live screen reaches: ' + ev(ws))
        return tag
    handed = [w for w in real if fclass(w['file']) == 'screen']
    if handed:
        lv = [w for w in handed if w['file'] in RFILES]
        if lv: return ('DERIVED', 'a screen builds it and hands it to a function: ' + ev(lv))
        return ('MISSING', 'built only on screen code no live route reaches: ' + ev(handed))
    return ('MISSING', 'no write found')

GENERIC = {'id', 'label', 'name', 'value', 'total', 'amount', 'date', 'type', 'kind', 'key', 'title', 'month', 'year', 'note', 'source',
           'status', 'text', 'sentence', 'reason', 'words', 'monthly', 'annual', 'rate', 'start', 'end', 'from', 'to', 'count', 'ok'}

inputs_rows = []
counted_inputs = 0
for fid in sorted(REGU):
    u = REGU[fid]; uid = u['uid']; hf = HF[fid]; fx = FACTS.get(fid, {}); nm = NAME_OF.get(fid, u['name'])
    mod = module_of(fid); seen_fields = set(); used = USE.get(fid, '')
    def add(inp, origin, evid, problem=''):
        inputs_rows.append({'Function id': fid, 'Function name': nm, 'Module': mod, 'Used?': used, 'Input': inp, 'Origin': origin, 'Evidence': evid, 'Problem': problem})
    # 1. registered inputs named on the header
    reads = re.findall(r'(MK-XS-I-\d{3})\s+([^\n]*)', hf.get('Reads', ''))
    for iid, iname in reads:
        info = INPUTS.get(iid, {}); s = (info.get('source') or '').lower()
        flds = [f.replace('\\', '') for f in INPUT_FIELDS.get(iid, [])]
        flds_clean = [re.sub(r'^[a-z]+\.', '', f).replace('b', '', 0) for f in flds]
        evid_parts = []; classes = set()
        for f in flds_clean:
            f2 = re.sub(r'[^A-Za-z0-9_]', '', f)
            if not f2: continue
            o, e = classify_field(f2)
            classes.add(o); evid_parts.append(f'{f2}: {o} — {e}')
        if 'typed' in s or 'set by the person' in s: decl = 'ENTERED'
        elif 'download' in s or 'connection' in s or 'feed' in s: decl = 'FEED'
        elif 'worked out' in s or 'the app' in s: decl = 'DERIVED'
        elif 'built in' in s: decl = 'DECIDED'
        elif 'person' in s and 'institution' in s: decl = 'ENTERED or FEED'
        else: decl = '?'
        if ('download' in s or 'connection' in s or 'feed' in s) and decl == 'ENTERED': decl = 'ENTERED or FEED'
        origin = decl
        problem = ''
        live_ok = classes & {'ENTERED', 'ENTERED or FEED', 'FEED', 'DERIVED'}
        if not live_ok:
            origin = 'MISSING'; problem = f'The register says "{info.get("source")}", but none of its fields ({", ".join(flds_clean)}) is set anywhere the live app reaches.'
        elif 'ENTERED' in decl and not any(c.startswith('ENTERED') for c in classes):
            problem = f'The register says a person types it ("{info.get("source")}"), but no live screen writes any of its fields — only {", ".join(sorted(classes))}.'
        add(f'{iid} ({iname.strip() or info.get("name")})', origin, f'register source: "{info.get("source")}". ' + ' | '.join(evid_parts[:4]), problem)
        counted_inputs += 1
        seen_fields |= set(re.sub(r'[^A-Za-z0-9_]', '', f) for f in flds_clean)
    # 2. parameters, and what live callers hand in
    live_calls = [c for c in CALLS if c['callee'] == fid and not c['test'] and c['from'] is not None and not is_demo(c['file'], c['from'])]
    live_calls_reached = [c for c in live_calls if c['from'] in REACHED]
    for pi, p in enumerate(fx.get('params', [])):
        props = p['props'] if p['props'] and p['type'].startswith('{') else []
        items = [(f"{p['name']}.{q}", q, pi) for q in props] if props else [(p['name'], None, pi)]
        for label, prop, idx in items:
            seen_fields.add(prop or p['name'])
            CLOCK = re.fullmatch(r'(now|today|asOf|nowMs|todayIso|at|date)', prop or p['name'] or '')
            if not live_calls:
                if used == 'USED':
                    add(label, 'DERIVED', f'parameter of {nm}; {nm} is handed around as a value (a callback or a table entry), so the compiler sees no direct call to read the argument from')
                else:
                    add(label, 'DERIVED', f'parameter of {nm}; no live code calls it ({used})')
                counted_inputs += 1; continue
            handed = []
            for c in (live_calls_reached or live_calls):
                if prop:
                    if prop in c['props']: handed.append((c, c['props'][prop]))
                    elif idx < len(c['args']) and not c['args'][idx].lstrip().startswith('{'):
                        handed.append((c, f"(carried inside `{c['args'][idx][:80]}`)"))
                    elif any(k.startswith('...') for k in c['props']): handed.append((c, '(spread: ' + ', '.join(v for k, v in c['props'].items() if k.startswith('...')) + ')'))
                else:
                    if idx < len(c['args']): handed.append((c, c['args'][idx]))
            n_calls = len(live_calls_reached or live_calls)
            optional = (prop and re.search(re.escape(prop) + r'\?\s*:', p['type'] or '')) or (not prop and p.get('optional'))
            if not handed and CLOCK:
                add(label, 'FEED', f'none of the {n_calls} live call(s) hands it in, so the function reads the device clock (e.g. {live_calls[0]["file"]}:{live_calls[0]["line"]})')
            elif not handed:
                add(label, 'MISSING', f'none of the {n_calls} live call(s) hands it in, e.g. {live_calls[0]["file"]}:{live_calls[0]["line"]}',
                    ('Optional input that no live caller supplies — the function always falls back to its own default.' if optional or p['name'] in ('opts', 'options', 'ctx') else
                     'Declared as an input but no live caller supplies it — whatever depends on it never runs.'))
            else:
                ex = handed[0]
                expr = ex[1]
                o = 'DERIVED'
                callee_ids = [cid for cid in re.findall(r'\b([A-Za-z_][A-Za-z0-9_]*)\(', expr) if cid in ISSUED]
                if re.fullmatch(r'-?[\d_.]+', expr.strip()): o = 'DECIDED?'
                elif re.search(r'\bstore\b|useStore|\bs\.\w', expr) or fclass(ex[0]['file']) == 'screen': o = 'from the screen (see field rows)'
                missing_calls = n_calls - len(handed)
                ev = f'{len(handed)} of {n_calls} live call(s) hand it in; e.g. {ex[0]["file"]}:{ex[0]["line"]} → `{expr[:140]}`'
                if callee_ids: ev += ' (made by ' + ', '.join(f'{ISSUED[c]} {c}' for c in callee_ids[:3]) + ')'
                prob = ''
                if o == 'DECIDED?': prob = f'A caller hands in a fixed number ({expr}); no decision is named for it.'
                if missing_calls and not optional: prob += f' {missing_calls} live call(s) leave it out.'
                add(label, 'DERIVED' if o != 'DECIDED?' else 'DECIDED?', ev, prob.strip())
            counted_inputs += 1
    # 3. stored fields it reads off its own inputs (not a function's output)
    for fld in fx.get('fields', []):
        f = fld['field']
        if f in seen_fields or fld.get('fromCall') is not None: continue
        if f in GENERIC: continue
        seen_fields.add(f)
        iid = FIELD_TO_INPUT.get(f)
        o, e = classify_field(f, uid)
        if o == 'DERIVED' and e.startswith('made inside this function'): continue   # a local the function builds itself
        if o == 'MISSING' and fclass(u['file']) == 'feed' and e.startswith('no code anywhere'):
            o, e = 'FEED', 'a field of the provider\'s own response or download, read by the feed code (our code never writes it; the provider does)'
        tag = f'{f}' + (f' (→ {iid})' if iid else ' (no input id)')
        prob = ''
        if o == 'MISSING': prob = 'Read here, but nothing the live app runs ever sets it — the branch that uses it is dead or always falls back.'
        add(tag, o, f'read at {u["file"]}:{fld["line"]} as `{fld["text"]}`; {e}', prob)
        counted_inputs += 1
    # 4. other functions' outputs
    callees = sorted({U[b]['id'] for b in OUTE[uid] if U[b].get('id') and U[b]['id'] != fid})
    for cid in callees:
        add(f'output of {cid} {NAME_OF.get(cid, "")}', 'DERIVED', f'{nm} calls it (compiler-resolved reference)', '' if USE.get(cid) != 'NOT IN CODE' else 'callee has no header')
        counted_inputs += 1
    # 5. fixed values: named constants
    for c in fx.get('consts', []):
        cname = c['name']
        if cname in ISSUED: continue   # a registered function written in capitals — counted as an output above
        lines = file_lines(c['file']); ctx = '\n'.join(lines[max(0, c['line'] - 6): c['line'] + 2]) + '\n' + c['value']
        if cname in ('DAY_MS', 'MS_PER_DAY', 'MONTHS', 'MONTH_NAMES'):
            add(f'constant {cname}', 'DECIDED', f'{c["file"]}:{c["line"]} — a calendar constant, not a money assumption'); counted_inputs += 1; continue
        has_num = bool(re.search(r'(?<![\w.])\d*\.?\d+(?![\w])', c['value'])) and not re.fullmatch(r"\s*'[^']*'\s*", c['value'])
        if cname in PUB:
            p = PUB[cname]; add(f'constant {cname}', 'PUBLISHED', f"{p['publication']} — valid through {p['valid_through']}, checked {p['checked']} ({c['file']}:{c['line']})")
        elif re.search(CITE_PUB, ctx) and has_num:
            add(f'constant {cname} = {c["value"][:60]}', 'PUBLISHED', f'{c["file"]}:{c["line"]} — the comment cites a published source: "{re.search(CITE_PUB, ctx).group(0)}"',
                'Not in the published-sources register (moneykeel_published_sources.json), so no expiry check covers it.')
        elif re.search(CITE_DEC, ctx):
            ds = re.findall(r'DEC-\d{4}', ctx)
            add(f'constant {cname} = {c["value"][:60]}', 'DECIDED', f'{c["file"]}:{c["line"]} — the comment cites: ' + (', '.join(ds) if ds else re.search(CITE_DEC, ctx).group(0)))
        else:
            ds = decs_naming(cname)
            if ds: add(f'constant {cname} = {c["value"][:60]}', 'DECIDED', f'{c["file"]}:{c["line"]} — named in the decision log: {", ".join(ds[:4])}')
            elif has_num and c.get('numeric', False) or (has_num and re.search(r'\d+\.\d+|\b[2-9]\d\b|\d{3,}', c['value'])):
                add(f'constant {cname} = {c["value"][:80]}', 'DECIDED?', f'{c["file"]}:{c["line"]} — no decision, published source or ruling cited, and the decision log does not name it',
                    'INVENTED NUMBER? A fixed number in code with no decision behind it.')
            else:
                add(f'constant {cname}', 'DECIDED', f'{c["file"]}:{c["line"]} — a fixed list of words or keys, not a number')
        counted_inputs += 1
    # 6. numbers written straight into the formula, and fallbacks used when an input is missing
    if 'Example only' in hf.get('Answers', '') or re.search(r'^(EXAMPLE|WORKED)|_EXAMPLE|EXAMPLE_', nm):
        add('its own numbers', 'DERIVED', 'worked example — the numbers in it are sample data for the figure register, not inputs the app uses')
        counted_inputs += 1
        continue
    num_seen = set()
    for d in fx.get('defaults', []):
        if d['num'] in ('0', '1', '0.0'): continue
        num_seen.add((d['num'], d['line']))
    for n in fx.get('nums', []) + [dict(x, ctx='') for x in fx.get('defaults', []) if x['num'] not in ('0', '1', '0.0')]:
        key = (n['num'], n['line'])
        if key in num_seen and n.get('ctx', None) == '' and any(z['line'] == n['line'] and z['num'] == n['num'] for z in fx.get('nums', [])): continue
        lines = file_lines(u['file']); ctx = n.get('ctx') or '\n'.join(lines[max(0, n['line'] - 5): n['line'] + 2])
        hdr_ctx = '\n'.join(ASSUME.get(nm, []))
        is_fallback = bool(re.search(r'\?\?\s*-?' + re.escape(n['num']) + r'\b|\|\|\s*-?' + re.escape(n['num']) + r'\b|(?<![<>=!])=\s*' + re.escape(n['num']) + r'\s*[,)]', n['text']))
        try: v = float(n['num'])
        except Exception: v = None
        pct = (f'{v * 100:g}%' if v is not None and 0 < v < 1 else None)
        mech = None
        if v is not None and (v <= 0.01 and v > 0) and re.search(r'Math\.abs|<=?|>=?|1e-', n['text']): mech = 'a rounding tolerance'
        elif re.search(r'for \(let \w+ = 0; \w+ < ' + re.escape(n['num']) + r'\b|while \(.*<\s*' + re.escape(n['num']) + r'\b', n['text']): mech = 'a loop cap'
        elif n['num'] in ('26', '52', '2080', '86400000', '1900', '36', '13', '2100', '9999') and not is_fallback: mech = 'a calendar or counting constant'
        dec_hit = []
        if v is not None:
            for d in DECS:
                if d.get('status') == 'SUPERSEDED': continue
                blob = ' '.join(str(d.get(k) or '') for k in ('plain', 'his_words', 'means', 'raw', 'governs', 'changed'))
                if nm in blob and (re.search(r'(?<![\d.])' + re.escape(n['num']) + r'(?![\d])', blob) or (pct and pct in blob)): dec_hit.append(d['id'])
        label = ('fallback ' if is_fallback else 'number ') + n['num'] + f' — `{n["text"][:110]}`'
        if re.search(CITE_PUB, ctx):
            add(label, 'PUBLISHED', f'{u["file"]}:{n["line"]} — the comment beside it cites "{re.search(CITE_PUB, ctx).group(0)}"')
        elif re.search(CITE_DEC, ctx):
            ds = re.findall(r'DEC-\d{4}', ctx)
            add(label, 'DECIDED', f'{u["file"]}:{n["line"]} — the comment beside it cites ' + (', '.join(ds) if ds else f'"{re.search(CITE_DEC, ctx).group(0)}"'))
        elif dec_hit:
            add(label, 'DECIDED', f'{u["file"]}:{n["line"]} — the decision log names this function with this number: {", ".join(dec_hit[:3])}')
        elif mech:
            add(label, 'DECIDED', f'{u["file"]}:{n["line"]} — {mech}, not a money assumption')
        elif n['num'] in hdr_ctx or (pct and pct in hdr_ctx):
            add(label, 'DECIDED', f'{u["file"]}:{n["line"]} — the function\'s recorded assumptions (moneykeel_code_formula ASSUMPTIONS_EXTRA) state it')
        else:
            add(label, 'DECIDED?', f'{u["file"]}:{n["line"]} — no decision, ruling or published source cited beside it',
                'INVENTED NUMBER? ' + ('A made-up value used when the real input is missing.' if is_fallback else 'A fixed number in the formula with no decision behind it.'))
        counted_inputs += 1

with_rows = {r['Function id'] for r in inputs_rows}
no_rows = sorted(set(REGU) - with_rows)
print(f'inputs rows: {len(inputs_rows)} across {len(with_rows)} of {len(REGU)} functions in the code; {len(no_rows)} read nothing at all: {no_rows[:12]}')
assert len(with_rows) + len(no_rows) == len(REGU) and len(inputs_rows) >= 3000
for fid in no_rows:
    inputs_rows.append({'Function id': fid, 'Function name': NAME_OF.get(fid, REGU[fid]['name']), 'Module': module_of(fid), 'Used?': USE.get(fid, ''),
                        'Input': '(none)', 'Origin': 'DERIVED', 'Evidence': 'takes no parameter, reads no field, constant or other function', 'Problem': ''})
print('origins:', collections.Counter(r['Origin'] for r in inputs_rows))

# ── inputs register: entered on a screen? read by a function? ──────────────────────────────────
reg_rows = []
for iid in sorted(INPUTS):
    info = INPUTS[iid]
    flds = [re.sub(r'[^A-Za-z0-9_]', '', re.sub(r'^[a-z]+\\?\.', '', f.replace('\\b', ''))) for f in INPUT_FIELDS.get(iid, [])]
    readers = sorted({U[x['unit']]['id'] for f in flds for x in READS.get(f, []) if x['unit'] is not None and U[x['unit']].get('id') and not is_test(x['file'])})
    origins = {f: classify_field(f) for f in flds if f}
    reg_rows.append({'Input': iid, 'Name': info['name'], 'Source (register)': info['source'], 'Fields': ', '.join(flds),
                     'Origin found': ' | '.join(f'{f}: {o[0]}' for f, o in origins.items()), 'Evidence': ' | '.join(f'{f}: {o[1]}' for f, o in origins.items())[:900],
                     'Read by functions': len(readers), 'Readers (first)': ', '.join(readers[:6])})

# ── entered on a screen, read by no function ───────────────────────────────────────────────────
entered = collections.defaultdict(list)
for c in STORECALLS:
    if fclass(c['file']) == 'screen' and c['file'] in RFILES:
        for k in c['keys']: entered[k].append(f"{c['file']}:{c['line']} ({c['action']})")
for k, ws in MENTIONS.items():
    for w in ws:
        if fclass(w['file']) == 'screen' and w['file'] in RFILES and w['file'].startswith('src/onboarding/'):
            t = file_lines(w['file'])[w['line'] - 1] if file_lines(w['file']) else ''
            if re.search(r"setAnswer\(\s*'" + re.escape(k) + "'|\\bk=\"" + re.escape(k) + '"|' + r"\bk:\s*'" + re.escape(k) + "'", t):
                entered[k].append(f"{w['file']}:{w['line']} (onboarding answer)")
for f in RFILES:
    if fclass(f) != 'screen': continue
    for m in re.finditer(r"setAnswer\(\s*'([A-Za-z_0-9]+)'|\bk=\"([A-Za-z_0-9]+)\"", src(f)):
        k = m.group(1) or m.group(2); entered[k].append(f"{f}:{src(f).count(chr(10), 0, m.start()) + 1} (answer)")
NOT_INPUT = {'id', 'label', 'name', 'updated_at', 'created_at', 'createdAt', 'ts', 'uid', 'positions', 'note', 'notes', 'source', 'xp', 'kind', 'type',
             'date', 'active', 'answers', 'bg', 'color', 'icon', 'hasError', 'stepIndex', 'tracks', 'receiptUri', 'cm', 'm', 'rt', 'sm', 'inviteCode'}
# keys a function reads through a name built at run time: a['ri_' + k], `${k}_freq`
DYN_PREFIX = set(); DYN_SUFFIX = set()
for f in BY_FILE:
    if is_test(f): continue
    for m in re.finditer(r"['\"]([A-Za-z]+_)['\"]\s*\+|`([A-Za-z]+_)\$\{", src(f)): DYN_PREFIX.add(m.group(1) or m.group(2))
    for m in re.finditer(r"\+\s*['\"](_[A-Za-z]+)['\"]|\}(_[A-Za-z]+)`", src(f)): DYN_SUFFIX.add(m.group(1) or m.group(2))
unused_rows = []
for k, where in sorted(entered.items()):
    if k in NOT_INPUT or len(k) < 3: continue
    wfiles = {w.split(':')[0] for w in where}
    rs = [x for x in READS.get(k, []) if not is_test(x['file'])] + [x for x in MENTIONS.get(k, []) if not is_test(x['file']) and x['file'] not in wfiles]
    by_fn = sorted({U[x['unit']]['id'] for x in rs if x['unit'] is not None and U[x['unit']].get('id')})
    by_calc = [x for x in rs if fclass(x['file']) in ('calc', 'store', 'feed')]
    by_other_screen = [x for x in rs if fclass(x['file']) == 'screen' and x['file'] not in wfiles]
    dyn = [p for p in DYN_PREFIX if k.startswith(p)] + [sfx for sfx in DYN_SUFFIX if k.endswith(sfx)]
    if by_fn or by_calc or dyn: continue
    if by_other_screen:
        status = f"UNSETTLED — no function reads it; another screen does ({by_other_screen[0]['file']}:{by_other_screen[0]['line']}) and may hand it to one"
    else:
        wlines = {(w.split(':')[0], w.split(':')[1].split(' ')[0]) for w in where}
        same = [x for x in READS.get(k, []) if x['file'] in wfiles and (x['file'], str(x['line'])) not in wlines and not is_test(x['file'])]
        if same:
            status = f"UNSETTLED — no function reads it; only the screen it is typed on reads it ({same[0]['file']}:{same[0]['line']}), which may use it there or only show it back in the box"
        else:
            status = 'ENTERED BUT UNUSED — nothing reads it except the box it is typed into'
    unused_rows.append({'Field': k, 'Entered at': '; '.join(sorted(set(where))[:3]), 'Read by': status})
print(f'entered fields: {len(entered)}; read by no registered function: {len(unused_rows)}')

json.dump({'use': use_rows, 'metrics': metric_rows, 'inputs': inputs_rows, 'register': reg_rows, 'entered_unused': unused_rows,
           'live': sorted(live), 'dead_routes': DEAD_ROUTES, 'door_of': door_of,
           'counts': {'functions': len(NAME_OF), 'in_code': len(REGU), 'metrics': len(METRICS), 'inputs_rows': len(inputs_rows),
                      'routes': len(ROUTES), 'live_routes': len(live), 'register_inputs': len(INPUTS)}},
          open(OUT, 'w'), default=str)
