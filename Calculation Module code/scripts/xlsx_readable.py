"""ONE WAY TO FORMAT A WORKBOOK, so every file he opens is readable.

The founder, 2026-09-27: *"in future learn how to format excel file so it is readable in terms of
column width."*

What was wrong, in three parts:

1. **Wrapped text with no row height.** Every cell was set to wrap and no row height was ever set, so
   a 300-character question rendered as a one-line sliver with the rest clipped — or, where Excel did
   auto-fit, as a row hundreds of points tall. Either way unreadable.
2. **Columns sized by guess.** The workplan's first two columns were 70 characters each: 140 characters
   before the STATUS column, which he therefore had to scroll sideways to reach.
3. **His columns last.** The orchestrator map put the question he must answer in column Q, behind
   thirteen columns of inventory that is mine to check, not his to read.

The rules here:
  · a column is as wide as its content needs, between MIN and MAX — measured, not guessed;
  · a row is as tall as its tallest wrapped cell needs, capped, so nothing is clipped and nothing is absurd;
  · the header row freezes, and carries a filter;
  · columns that are MINE can be grouped and collapsed, so the file opens on his columns.

Call `readable(ws, ...)` after the rows are written.
"""
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill

MIN_W, MAX_W = 9, 52          # characters. Past ~52 the eye loses the line start on the way back.
LINE_H = 15                   # points per line of 11-point text
MAX_LINES = 4                 # a cell taller than this is read by clicking it, not by scrolling past it
# WHY FOUR, not eight. Eight lines is 120 points — about five rows on a laptop screen, so a 189-row tab
# becomes forty screens of scrolling. Four lines shows the start of every cell and the whole of most of
# them; anything longer is read in the formula bar by clicking the cell. A tab whose whole purpose is a
# paragraph he must READ IN PLACE — the Read me, the question he answers — passes max_lines itself.
HEAD = Font(bold=True, color='FFFFFF', size=11)
FILL = PatternFill('solid', fgColor='1F4E78')


def _lines(text: str, width: int) -> int:
    """How many lines this cell needs at that column width, counting its own line breaks."""
    n = 0
    for para in str(text).split('\n'):
        n += max(1, -(-len(para) // max(1, width - 2)))     # ceil
    return n


def readable(ws, *, widths=None, mine=None, freeze='A2', filt=True, max_lines=MAX_LINES):
    """Size every column and row of `ws` so a person can read it.

    widths    optional {column letter: width} to pin a column (a column he types in wants room).
    mine      optional [column letters] that are generated inventory — grouped and collapsed, so the
              sheet opens on the columns he actually reads.
    """
    widths = widths or {}
    ncols, nrows = ws.max_column, ws.max_row
    if not nrows or not ncols:
        return ws

    # ── columns: the widest of the header and the longest single LINE in the column, capped
    for c in range(1, ncols + 1):
        letter = get_column_letter(c)
        if letter in widths:
            ws.column_dimensions[letter].width = widths[letter]
            continue
        longest = 0
        for r in range(1, nrows + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            for para in str(v).split('\n'):
                longest = max(longest, len(para))
        ws.column_dimensions[letter].width = max(MIN_W, min(MAX_W, longest + 2))

    # ── rows: tall enough for the tallest wrapped cell in the row, and never taller than the cap.
    # A FOLDED COLUMN DOES NOT COUNT (2026-09-27): the Map's inventory columns hold hundred-line lists,
    # and letting them drive the height made every row nine lines tall to fit text nobody can see.
    folded = set(mine or [])
    for r in range(1, nrows + 1):
        need = 1
        for c in range(1, ncols + 1):
            letter = get_column_letter(c)
            if letter in folded:
                continue
            v = ws.cell(r, c).value
            if v is None:
                continue
            w = ws.column_dimensions[letter].width or MAX_W
            need = max(need, _lines(v, int(w)))
        ws.row_dimensions[r].height = min(need, max_lines) * LINE_H

    # ── the header: legible, frozen, and filterable
    for c in ws[1]:
        c.font = HEAD
        c.fill = FILL
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = freeze
    if filt and nrows > 1:
        ws.auto_filter.ref = f'A1:{get_column_letter(ncols)}{nrows}'

    # ── my columns, folded away
    for letter in (mine or []):
        ws.column_dimensions[letter].outlineLevel = 1
        ws.column_dimensions[letter].hidden = True
    if mine:
        ws.sheet_properties.outlinePr.summaryRight = False
    return ws
