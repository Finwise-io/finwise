# Converts the workflow's decomposition files into ONE reviewable spreadsheet.
#
# Mechanical only: the words come from the drafting agents and their critics, never from here.
# Writes a STANDALONE preview file — nothing merges into the master until the founder has read
# the open questions and given verdicts. That gate is his, not a build step.
#
# Expected format per section file (strict, defined in the workflow brief):
#   ### <Metric>            — one line under it: what the person is really asking
#   #### P<n> — <purpose>
#   | Condition | Case | Calculation or refusal |   (rows A missing … F account-type)
#   ## Open questions for the founder
import re, os, sys, glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = "/Users/palakjain/Vibe_Coding/finwise/Calculation Module"
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "decomposition")
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, "decomposition", "use-case-grid-PREVIEW.xlsx")

SECTION_TITLES = {"A": "Net worth & balance sheet", "B": "Portfolio performance",
                  "C": "Budget, spending & cash flow", "D": "Retirement & planning",
                  "E": "Cross-cutting rules"}

rows = []          # (section, metric, ask, purpose, cond, case, rule)
questions = []     # (section, text)
parse_errors = []

for path in sorted(glob.glob(os.path.join(SRC, "section-*.md"))):
    sec = os.path.basename(path).split("-")[1].split(".")[0].upper()
    text = open(path, encoding="utf-8").read()
    metric = ask = purpose = None
    lines = text.split("\n")
    i = 0
    in_questions = False
    while i < len(lines):
        ln = lines[i].rstrip()
        if ln.startswith("## Open questions"):
            in_questions = True; i += 1; continue
        if in_questions:
            m = re.match(r"\s*\d+[.)]\s+(.*)", ln)
            if m: questions.append((sec, m.group(1).strip()))
            i += 1; continue
        if ln.startswith("### ") and not ln.startswith("####"):
            metric = ln[4:].strip(); ask = None; purpose = None
            # the next non-empty line is the "really asking" line
            j = i + 1
            while j < len(lines) and not lines[j].strip(): j += 1
            if j < len(lines) and not lines[j].startswith("#"):
                ask = lines[j].strip()
            i += 1; continue
        if ln.startswith("#### "):
            purpose = ln[5:].strip(); i += 1; continue
        if ln.startswith("|") and metric and purpose:
            cells = [c.strip() for c in ln.strip("|").split("|")]
            if len(cells) >= 3 and cells[0] not in ("Condition", "---", ":---", "") \
               and not set(cells[0]) <= set("-: "):
                rows.append((sec, metric, ask or "", purpose, cells[0], cells[1], cells[2]))
        i += 1
    if not any(r[0] == sec for r in rows):
        parse_errors.append(f"{os.path.basename(path)}: no grid rows parsed")

if not rows:
    sys.exit("nothing parsed — are the section files present?")

wb = openpyxl.Workbook()
HDR = PatternFill("solid", fgColor="085041"); WHITE = Font(color="FFFFFF", bold=True, size=10)
WRAP = Alignment(wrap_text=True, vertical="top"); THIN = Border(*[Side(style="thin", color="DDDDDD")] * 4)
SECF = PatternFill("solid", fgColor="D9E2EC"); METF = PatternFill("solid", fgColor="EDF2F0")
REF  = PatternFill("solid", fgColor="FBE3E3"); YOURS = PatternFill("solid", fgColor="E8F0FB")

ws = wb.active; ws.title = "Use-case grid"
ws.sheet_properties.tabColor = "085041"
ws.merge_cells("A1:H1")
ws.cell(1, 1, "USE-CASE GRID — every metric × its purposes × the six data conditions. Drafted and "
              "adversarially verified by the workflow's agents; converted mechanically; no cell "
              "written by Claude Code. Red cells are refusals (the honest 'we do not show a number "
              "here'). PREVIEW — merges into the master only after your verdicts."
).font = Font(bold=True, size=11, color="085041")
ws.cell(1, 1).alignment = WRAP; ws.row_dimensions[1].height = 46
COLS = [("Section", 10), ("Metric", 24), ("Purpose", 30), ("Condition", 16),
        ("Case", 34), ("Calculation or refusal", 66), ("YOUR VERDICT", 14), ("YOUR COMMENT", 30)]
for c, (t, w) in enumerate(COLS, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    cell = ws.cell(2, c, t); cell.fill, cell.font, cell.alignment, cell.border = HDR, WHITE, WRAP, THIN
ws.freeze_panes = "C3"

r = 3; last_sec = last_met = None
REFUSAL = re.compile(r"refus|not shown|no figure|absent|suppress|we do not|don't show|do not show", re.I)
for sec, metric, ask, purpose, cond, case, rule in rows:
    if sec != last_sec:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        cell = ws.cell(r, 1, f"{sec} · {SECTION_TITLES.get(sec, sec)}")
        cell.fill, cell.font, cell.border = SECF, Font(bold=True, size=10, color="1B3A5C"), THIN
        last_sec, last_met = sec, None; r += 1
    if metric != last_met:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        cell = ws.cell(r, 1, f"{metric} — {ask}")
        cell.fill, cell.font, cell.alignment, cell.border = METF, Font(bold=True, size=10), WRAP, THIN
        last_met = metric; r += 1
    vals = [sec, metric, purpose, cond, case, rule, "", ""]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(r, c, v); cell.alignment, cell.border = WRAP, THIN
        if c == 6 and REFUSAL.search(str(v)): cell.fill = REF
        if c in (7, 8): cell.fill = YOURS
    r += 1

qs = wb.create_sheet("Open questions")
qs.sheet_properties.tabColor = "A32D2D"
qs.cell(1, 1, "Open questions the agents could not decide as the CFA — your call, one per row."
).font = Font(bold=True, size=11, color="A32D2D")
for c, (t, w) in enumerate([("#", 5), ("Section", 9), ("Question", 90), ("Your answer", 40)], start=1):
    qs.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    cell = qs.cell(2, c, t); cell.fill, cell.font, cell.border = HDR, WHITE, THIN
oq = os.path.join(SRC, "OPEN-QUESTIONS.md")
qrows = questions
if os.path.exists(oq):
    qrows = [("", re.match(r"\s*\d+[.)]\s*(.*)", l).group(1))
             for l in open(oq, encoding="utf-8") if re.match(r"\s*\d+[.)]\s", l)]
for i, (sec, q) in enumerate(qrows, start=1):
    qs.cell(i + 2, 1, i); qs.cell(i + 2, 2, sec); qs.cell(i + 2, 3, q)
    qs.cell(i + 2, 3).alignment = WRAP
    qs.cell(i + 2, 4).fill = YOURS

wb.save(OUT)
mets = len({(s, m) for s, m, *_ in rows})
purps = len({(s, m, p) for s, m, a, p, *_ in rows})
print(f"{OUT}")
print(f"  {mets} metrics · {purps} purpose cases · {len(rows)} condition cells · {len(qrows)} open questions")
if parse_errors:
    print("  ⚠ " + " · ".join(parse_errors))
