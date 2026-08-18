# Builds the "Raw inputs" tab — every value that enters the app from OUTSIDE, and what we call it.
#
# Founder, 2026-08-17: "is it feasible to have one tab detailing every raw input we are capturing
# and source - and if we have our name for that input."
#
# The definition that makes the tab meaningful: a RAW input is one that is not computed from
# anything else. Somebody typed it, an institution sent it, a feed returned it, or WE chose it.
# Everything else in the inventory is derived, and its row says what it is derived from.
#
# The last category matters most and is the easiest to leave out: numbers WE chose. A long-run
# return of 10.4% for equities is not a fact an institution sent us — it is an assumption, and
# every projection stands on it. Listing it beside real inputs is the honest thing to do.
import os, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = "/Users/palakjain/Vibe_Coding/finwise"
OUT  = os.path.join(ROOT, "Calculation Module/calculation-inventory-v1.xlsx")

HDR  = PatternFill("solid", fgColor="085041")
WHITE= Font(color="FFFFFF", bold=True, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(*[Side(style="thin", color="DDDDDD")] * 4)
YOURS= PatternFill("solid", fgColor="E8F0FB")
BANDS = {
  "TYPED":  PatternFill("solid", fgColor="FFF6E5"),
  "BROKER": PatternFill("solid", fgColor="EAF3EC"),
  "FILE":   PatternFill("solid", fgColor="EAF0F8"),
  "FEED":   PatternFill("solid", fgColor="F5EDF8"),
  "OURS":   PatternFill("solid", fgColor="FBE9E9"),
}
SECTION = PatternFill("solid", fgColor="D9E2EC")

# (band, what it is, where it comes from, their name for it, our name, used by)
R = [
 ("TYPED", "SECTION", "TYPED BY THE PERSON — nothing external supplies these, so they are only as right as the question that asked for them", "", "", ""),
 ("TYPED", "What an account is called", "The Name box on Add account; a connected account arrives pre-named by the broker and can be renamed",
  "Account name / nickname", "label",
  "ASSET_CLASS_OF_ACCOUNT · CLASS_FROM_NAME_TEXT — the name is READ as data, so renaming an account can change its class"),
 ("TYPED", "Which institution holds it", "The Institution box, or the broker at connect time", "Institution", "institution", "IMPORT_ACCOUNT_MATCH · account grouping on Net worth"),
 ("TYPED", "What kind of account it is", "Picked from the list on Add account — Checking, Savings, High-yield savings, Money market, CD, Stocks & Funds, Fixed income, Crypto, Options, Home, Vehicle…",
  "(no institution supplies this)", "kind", "ASSET_CLASS_OF_ACCOUNT · EXPECTED_RETURN_BY_KIND · KIND_IS_CASH · EARMARK_DEFAULT"),
 ("TYPED", "The wrapper it sits in, for tax", "The wrapper picked on Add account — Taxable / 401(k) / Traditional IRA / Roth / HSA",
  "Account type", "tax_bucket", "TAX_TREATMENT · ASSET_CLASS_OF_ACCOUNT (last resort) · the tax split in the retirement engine"),
 ("TYPED", "What it is worth", "The Balance box — for a HAND-ENTERED account. A connected account's balance comes from the broker instead",
  "Balance / Value", "balance", "every total: TOTAL_CASH · TOTAL_INVESTMENTS · TOTAL_INVESTABLE · ALLOCATION_BY_CLASS · NEST_EGG"),
 ("TYPED", "When that value was last true", "Set when a hand-entered value is saved or confirmed", "(ours)", "value_as_of", "the freshness line on Net worth and Account detail"),
 ("TYPED", "How much of it is for retirement", "The per-account slider on the Plan tab. Absent means use the default rule",
  "(ours)", "retirement_pct", "EARMARK_AMOUNT · NEST_EGG · EXPECTED_RETURN_BLENDED"),
 ("TYPED", "Which class it is, when overridden", "Chosen on the account screen — always beats every derived rule",
  "(ours)", "asset_class", "ASSET_CLASS_OF_ACCOUNT — rule 1, the person's own answer"),
 ("TYPED", "How a home is used", "Primary / rental / second home / land, on the property screen", "(ours)", "re_use", "EARMARK_DEFAULT — only a PRIMARY home is excluded from the nest egg"),
 ("TYPED", "A bond's face value, coupon and maturity", "The Add bond form — or parsed out of a brokered CD's description text",
  "Face / Par · Coupon · Maturity", "face_value · coupon_rate · maturity_date", "BOND_COUPON · BOND_CURRENT_YIELD · BOND_YIELD_TO_MATURITY · BOND_YEARS_LEFT · MATURITY_MEANS_BOND"),
 ("TYPED", "Every debt", "The Debts screen. NO institution supplies debts today — all of it is typed",
  "(none)", "remaining_balance · interest_rate_apr · minimum_monthly_payment · monthly_payment · payment_type · debt_type",
  "DEBT_REQUIRED_PAYMENT · DEBT_MINIMUM_TOTAL · DEBT_ACTUAL_TOTAL · LOAN_PAYMENT · MONTHS_TO_CLEAR"),
 ("TYPED", "A credit limit", "The card form. WHEN MISSING, utilisation currently reads as 0%, which looks like perfect behaviour",
  "Credit limit", "credit_limit", "CREDIT_UTILISATION — see its DEFECT row"),

 ("BROKER", "SECTION", "SENT BY A CONNECTED BROKER (SnapTrade) — the exact payload paths, as the live E*TRADE and Vanguard connections deliver them", "", "", ""),
 ("BROKER", "What the account is worth", "accounts[].balance.total.amount", "Total value", "balance", "every total on Net worth"),
 ("BROKER", "Uninvested cash in the account", "accounts[].balances.cash — WARNING: this INCLUDES money-market funds that also appear as holdings",
  "Cash / Sweep", "cash_balance", "CASH_SLEEVE, which subtracts the money-market funds so the money is not counted twice"),
 ("BROKER", "The ticker of a holding", "positions[].symbol.raw_symbol (falling back to .symbol). A brokered CD arrives with a CUSIP here, not a ticker",
  "Symbol", "ticker", "CLASS_FROM_BROKER · TICKER_PLAUSIBLE · every price lookup"),
 ("BROKER", "The name of a holding", "positions[].symbol.description — THE most important text field we receive: for a CD or a bond it is the ONLY place the maturity and coupon appear",
  "Description / Security name", "name / label", "CLASS_FROM_BROKER · BOND_YEARS_LEFT · BOND_COUPON"),
 ("BROKER", "The broker's own type for it", "positions[].symbol.type.code — cs (common share) · et (ETF) · mf (mutual fund) · oef/cef · bnd (bond) · crypto",
  "Security type", "(read directly, not stored)", "CLASS_FROM_BROKER — but Vanguard sends brokered CDs with an ORDINARY EQUITY code, so the description is trusted over this"),
 ("BROKER", "How many shares", "positions[].units", "Quantity", "shares", "HOLDING_SHARES · HOLDING_MARKET_VALUE"),
 ("BROKER", "The broker's price", "positions[].price — as fresh as the last sync, which is why prices are fetched separately too",
  "Price / Last", "last_price", "HOLDING_MARKET_VALUE where no live price exists"),
 ("BROKER", "What was paid per share", "positions[].average_purchase_price", "Average cost", "cost_per_share", "HOLDING_COST_BASIS · HOLDING_UNREALISED_GAIN"),
 ("BROKER", "Purchase history", "positions[].tax_lots[] — original_purchase_date, quantity, purchased_price. MOST BROKERS SEND NONE, and then ONE synthetic lot is built from the average price",
  "Tax lots / Purchase history", "lots[]", "HOLDING_CAP_GAINS (long vs short term) · HOLDING_FIRST_BOUGHT"),
 ("BROKER", "Whether it is treated as cash", "positions[].cash_equivalent", "Cash equivalent", "cashEquivalent", "CASH_SLEEVE — and the F6 defect, where the fund is removed from holdings as well"),
 ("BROKER", "Every activity", "activities[] — type, date, amount, units, price, and the broker's own id",
  "Transactions / Activity / History", "date · type · amount · shares · price", "INGEST_ACTIVITY_TYPE · INGEST_DEDUPE · NW_CHANGE_WALK · the money-weighted return"),
 ("BROKER", "How far back the history goes", "The earliest activity the source has ever supplied. Brokers differ — one sends a year, the next 90 days",
  "(not stated by the broker; we observe it)", "history_from", "year-scale income figures label themselves with it, so a short history cannot silently understate"),
 ("BROKER", "Option positions", "option_positions[].symbol.option_symbol — underlying, strike, expiry, call/put",
  "Options", "option_holdings[]", "shown as itemised detail; their value is already inside the broker's total and is never added again"),

 ("FILE", "SECTION", "READ FROM A FILE YOU UPLOAD — no two institutions name a column the same way, so every column is found by PATTERN and the importer reports which header it used", "", "", ""),
 ("FILE", "The ticker", "A header matching Symbol or Ticker", "Symbol · Ticker", "symbol / ticker", "CLASS_FROM_FILE · FILE_IMPORT"),
 ("FILE", "The name of the holding", "Description · Security · Name · Fund · Investment · Holding · Instrument", "(all six seen in the wild)", "label / name", "CLASS_FROM_FILE — the words in this column decide the class"),
 ("FILE", "How many shares", "Quantity · Shares · Qty · Units", "", "shares", "HOLDING_SHARES"),
 ("FILE", "What it is worth", "Value · Market Value · Current Value", "", "value", "the account total"),
 ("FILE", "What was paid", "Cost Per Share · Cost/Sh · Average Cost · Price Paid · Purchase Price · Unit Cost — falling back to Cost Basis · Total Cost · Book Cost divided by the quantity",
  "", "costPerShare", "HOLDING_COST_BASIS · HOLDING_UNREALISED_GAIN"),
 ("FILE", "When it was bought", "Purchase Date · Acquired · Acquisition · Trade Date · Date. Read as YYYY-MM-DD or MM/DD/YYYY; anything else is left blank rather than guessed",
  "", "date", "HOLDING_CAP_GAINS · HOLDING_FIRST_BOUGHT"),

 ("FEED", "SECTION", "FETCHED FROM AN OUTSIDE FEED — not from your institution, and each carries its own freshness", "", "", ""),
 ("FEED", "Share prices", "The market-data service (Tiingo / Yahoo), keyed by ticker, cached with the time of the last successful fetch",
  "Daily close", "points[].close", "PRICE_LATEST · HOLDING_MARKET_VALUE · PERIOD_RETURN · PRICE_FRESHNESS"),
 ("FEED", "The 10-year Treasury yield", "FRED (the Federal Reserve's own series, DGS10). The previous endpoint returned 404 and was silently dead",
  "DGS10", "treasuryYield", "the mortgage-versus-Treasury comparison. Falls back to a stated figure, and says so when it does"),
 ("FEED", "Inflation", "The published series, with a stated fallback when it cannot be reached", "CPI", "inflation", "the retirement engine · cost-of-living adjustments"),

 ("OURS", "SECTION", "NUMBERS WE CHOSE — no institution sent these. They are assumptions, every projection stands on them, and they belong on this tab precisely because they look like facts", "", "", ""),
 ("OURS", "The long-run return for each account type", "The ASSET_KINDS table in our code: checking 0.5% · savings 2.4% · high-yield 4.2% · money market 4.5% · CD 4.2% · stocks and funds 10.4% …",
  "(ours)", "ASSET_KINDS[].ret", "EXPECTED_RETURN_BY_KIND · EXPECTED_RETURN_BLENDED — and through those, every projection and every odds figure"),
 ("OURS", "How much a year's return typically swings", "The volatility figure per account type, added 2026-08-11 to replace a rule of thumb (return × 1.7) that made a cash-heavy pot look almost risk-free",
  "(ours)", "ASSET_KINDS[].vol", "EXPECTED_SWING_BLENDED · the spread of the simulation"),
 ("OURS", "Which heading each account type sits under", "The bucket and section in ASSET_KINDS", "(ours)", "ASSET_KINDS[].bucket · .section",
  "the Net worth grouping — and the F6 contradiction, where money market is marked CASH here while its class says otherwise"),
 ("OURS", "Capital-gains rates", "A stated default where the person has not given their own bracket", "(ours)", "long / short rates", "CAP_GAINS_TAX"),
 ("OURS", "Which fund each holding is compared against", "Our mapping from account type to a benchmark ticker", "(ours)", "benchmarkTicker", "PERFORMANCE — the beat-by column"),
 ("OURS", "Where the lines fall", "Stale after 4 days · utilisation good below 30% and caution below 50% · concentration reported at 25% or more",
  "(ours)", "the thresholds", "PRICE_FRESHNESS · CREDIT_UTILISATION · the concentration card"),
]

wb = openpyxl.load_workbook(OUT)
# Carry the founder's column and any cell comment across the rebuild, keyed on "What it is".
keep, keepNotes = {}, {}
if "Raw inputs" in wb.sheetnames:
    old = wb["Raw inputs"]
    for _r in range(3, old.max_row + 1):
        _k = old.cell(_r, 1).value
        if _k and old.cell(_r, 6).value: keep[_k] = old.cell(_r, 6).value
        for _c in range(1, 7):
            _cell = old.cell(_r, _c)
            if _k and _cell.comment: keepNotes[(_k, _c)] = (_cell.comment.text, _cell.comment.author)
    del wb["Raw inputs"]
ws = wb.create_sheet("Raw inputs", 2)
ws.sheet_properties.tabColor = "1B4D8F"
ws.merge_cells("A1:F1")
ws.cell(1, 1, "RAW INPUTS — every value that enters the app from OUTSIDE, and what we call it. "
              "A raw input is one that is not computed from anything else: somebody typed it, an institution sent it, "
              "a feed returned it, or WE chose it. Everything else in this workbook is derived, and its own row says "
              "what it is derived from. READ THE LAST SECTION CAREFULLY — those are numbers we chose, not facts anyone "
              "sent us, and every projection stands on them.").font = Font(bold=True, size=11, color="1B4D8F")
ws.cell(1, 1).alignment = WRAP
ws.row_dimensions[1].height = 60

COLS = [("What it is", 34), ("Where it comes from — the exact door and field", 62),
        ("Their name for it", 26), ("OUR name for it", 30),
        ("What uses it", 52), ("YOUR COMMENT / QUESTION", 40)]
for c, (t, w) in enumerate(COLS, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    cell = ws.cell(2, c, t); cell.fill, cell.font, cell.alignment, cell.border = HDR, WHITE, WRAP, THIN

r = 3
for band, what, where, theirs, ours, used in R:
    if what == "SECTION":
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        cell = ws.cell(r, 1, where)
        cell.fill, cell.font, cell.alignment, cell.border = SECTION, Font(bold=True, size=10, color="1B3A5C"), WRAP, THIN
        ws.row_dimensions[r].height = 30
    else:
        for c, v in enumerate([what, where, theirs, ours, used, ""], start=1):
            cell = ws.cell(r, c, v)
            cell.alignment, cell.border = WRAP, THIN
            cell.fill = YOURS if c == 6 else BANDS[band]
            if c == 6 and what in keep: cell.value = keep[what]
            if (what, c) in keepNotes:
                _t, _a = keepNotes[(what, c)]
                cell.comment = openpyxl.comments.Comment(_t, _a or "Founder")
    r += 1
ws.freeze_panes = "A3"
wb.save(OUT)
print(f"Raw inputs tab: {len([x for x in R if x[1] != 'SECTION'])} inputs across "
      f"{len([x for x in R if x[1] == 'SECTION'])} sources")
