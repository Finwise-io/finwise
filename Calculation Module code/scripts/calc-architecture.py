# The DETAILED Calculation Module architecture — every engine, its layer, what feeds it, its status.
# Generated from the inventory workbook, never drawn by hand.
import openpyxl, json, collections, re

P='Calculation Module/calculation-inventory-v1.xlsx'
ws=openpyxl.load_workbook(P)['Calculations']
h={ws.cell(2,c).value:c for c in range(1,ws.max_column+1)}

def part(src):
    s=str(src or '')
    for pat,name in [('domain/retirement','Plan & retirement'),('domain/planning','Plan & retirement'),
        ('domain/decumulation','Plan & retirement'),('domain/performance','Performance'),
        ('domain/bonds','Performance'),('domain/alternatives','Performance'),
        ('domain/assets','Net worth'),('domain/networth','Net worth'),('domain/snapshot','Net worth'),
        ('domain/history','Net worth'),('domain/income','Cash flow'),('domain/budget','Cash flow'),
        ('domain/paycheck','Cash flow'),('domain/transactions','Cash flow'),('domain/cashflow','Cash flow'),
        ('domain/debt','Debt'),('domain/goals','Goals'),('domain/import','Accounts & connections'),
        ('services/sync','Accounts & connections'),('services/marketData','Market data'),
        ('services/economicData','Market data'),('domain/insights','Insights'),('domain/gaps','Insights'),
        ('domain/profile','Setup'),('onboarding','Setup'),('domain/_shared','Foundation')]:
        if pat in s: return name
    return 'Other'

fn2eng={}
rows=[]
for r in range(3,ws.max_row+1):
    eng=ws.cell(r,1).value
    if not eng: continue
    fn=ws.cell(r,2).value
    fn2eng[fn]=eng
    lay=str(ws.cell(r,h['Build order']).value or '1')
    m=re.match(r'(\d+)',lay)
    rows.append(dict(engine=eng, fn=fn,
        layer=int(m.group(1)) if m else 1,
        builtFrom=str(ws.cell(r,h['Built from']).value or '—'),
        status=ws.cell(r,h['Status']).value or 'PINNED',
        part=part(ws.cell(r,h['Source']).value),
        screens=str(ws.cell(r,h['Screens']).value or '—'),
        answers=(str(ws.cell(r,3).value or '').split('\n')[0])[:110],
        probe=str(ws.cell(r,h['Probe']).value or ''),
        source=str(ws.cell(r,h['Source']).value or '')))

# translate builtFrom function names into engine names
for x in rows:
    deps=[d.strip() for d in x['builtFrom'].split(',') if d.strip() and d['builtFrom' and 0:] is not None]
    deps=[d.strip() for d in x['builtFrom'].split(',') if d.strip() and d.strip()!='—']
    x['feeds_from']=[fn2eng.get(d,d) for d in deps]
    del x['builtFrom']

json.dump({'generated':'2026-08-18','engines':rows}, open('docs/architecture/calc-module-map.json','w'), indent=1)
layers=collections.Counter(x['layer'] for x in rows)
parts=collections.Counter(x['part'] for x in rows)
print(f"{len(rows)} engines · layers: " + ' '.join(f"L{k}:{v}" for k,v in sorted(layers.items())))
print("parts: " + ' · '.join(f"{k} {v}" for k,v in parts.most_common()))
