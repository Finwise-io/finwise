# Assembles the contracted deliverable — MECHANICAL ONLY. Every word comes from the workflow's
# verified files, the spine, the master, or the field-crosswalk research. Nothing is authored here.
#
# THE CONTRACT (founder, 2026-08-19):
#   Tab 1 grid: metrics · use cases · CURRENT formula · NEW formula · variables/inputs · sources —
#     and the new formula is PLAIN TEXT IN LITERALLY COLUMN K.
#   Tab 2: every variable — is it provided by SnapTrade / Yodlee / Plaid, and the potential names.
#   Rule: nothing reaches an engine that is not written in this file.
import re, os, sys, glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = "/Users/palakjain/Vibe_Coding/finwise/Calculation Module"
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "decomposition")
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(SRC, "METRICS-DELIVERABLE-PREVIEW.xlsx")

HDR = PatternFill("solid", fgColor="085041"); WHITE = Font(color="FFFFFF", bold=True, size=10)
WRAP = Alignment(wrap_text=True, vertical="top"); THIN = Border(*[Side(style="thin", color="DDDDDD")] * 4)
SECF = PatternFill("solid", fgColor="D9E2EC"); METF = PatternFill("solid", fgColor="EDF2F0")
REF = PatternFill("solid", fgColor="FBE3E3"); YOURS = PatternFill("solid", fgColor="E8F0FB")
GEN = PatternFill("solid", fgColor="EEEEE9"); KFILL = PatternFill("solid", fgColor="FFF9E8")
SECTION_TITLES = {"A": "Net worth & balance sheet", "B": "Portfolio performance",
                  "C": "Budget, spending & cash flow", "D": "Retirement & planning",
                  "E": "Cross-cutting rules"}

# ── joins: spine (inputs, trap, availability) · master (current formula, variable sources) ──────
spine = openpyxl.load_workbook(os.path.join(ROOT, "investment-api-performance-crosswalk-Aug19.xlsx"))
t9 = spine["9. Comprehensive Metrics"]
SPINE = {}
for r in range(5, t9.max_row + 1):
    if t9.cell(r, 1).value and t9.cell(r, 2).value:
        SPINE[str(t9.cell(r, 1).value).strip().lower()] = {
            "inputs": t9.cell(r, 5).value or "", "avail": t9.cell(r, 6).value or "",
            "trap": t9.cell(r, 7).value or ""}

master = openpyxl.load_workbook(os.path.join(ROOT, "MoneyKeel-calculation-master-v1.xlsx"))
m1 = master["1. Metrics"]; m2 = master["2. Engines"]
h2 = {m2.cell(2, c).value: c for c in range(1, m2.max_column + 1)}
ENG = {}
for r in range(3, m2.max_row + 1):
    name = m2.cell(r, h2["Engine name"]).value
    if name:
        ENG[str(name)] = {
            "formula": m2.cell(r, h2["Formula — the actual rules"]).value or "",
            "vars": m2.cell(r, h2["Input variables (raw → ours)"]).value or "",
            "src": m2.cell(r, h2["Where each comes from — file, feed, field"]).value or "",
            "cite": m2.cell(r, h2.get("Source", 0)).value or ""}
MET2ENG = {}
for r in range(5, m1.max_row + 1):
    met = m1.cell(r, 1).value; cited = str(m1.cell(r, 8).value or "")
    if met and cited:
        names = [t for t in re.split(r"[·/,()\s]+", cited) if t.isupper() and "_" in t]
        if names: MET2ENG[str(met).strip().lower()] = names

def current_of(metric):
    engs = MET2ENG.get(metric.strip().lower(), [])
    parts = []
    for e in engs[:2]:
        if e in ENG and ENG[e]["formula"]:
            f = str(ENG[e]["formula"]).split("\n")[0]
            parts.append(f"[{e}] {f}")
    return "\n".join(parts) if parts else "(not built today)"

def vars_of(metric):
    s = SPINE.get(metric.strip().lower(), {})
    engs = MET2ENG.get(metric.strip().lower(), [])
    ev = "\n".join(str(ENG[e]["vars"]).strip() for e in engs[:1] if e in ENG and ENG[e]["vars"])
    return (str(s.get("inputs", "")).strip() + ("\n— engine today: " + ev if ev else "")).strip() or "(see K)"

def src_of(metric):
    engs = MET2ENG.get(metric.strip().lower(), [])
    for e in engs:
        if e in ENG and ENG[e]["src"]:
            return str(ENG[e]["src"]).split("\n\n")[0][:400]
    return "(per variable — see tab 2)"

# ── parse the verified section files ─────────────────────────────────────────────────────────────
rows = []
for path in sorted(glob.glob(os.path.join(SRC, "section-*.md"))):
    sec = os.path.basename(path).split("-")[1].split(".")[0].upper()
    metric = purpose = None; in_q = False
    for ln in open(path, encoding="utf-8"):
        ln = ln.rstrip()
        if ln.startswith("## Open questions"): in_q = True
        if in_q: continue
        if ln.startswith("### ") and not ln.startswith("####"):
            metric = ln[4:].strip(); purpose = None
        elif ln.startswith("#### "):
            purpose = ln[5:].strip()
        elif ln.startswith("|") and metric and purpose:
            cells = [c.strip() for c in ln.strip("|").split("|")]
            if len(cells) >= 3 and cells[0] not in ("Condition", "") and not set(cells[0]) <= set("-: "):
                rows.append((sec, metric, purpose, cells[0], cells[1], cells[2]))
if not rows: sys.exit("no grid rows parsed")

wb = openpyxl.Workbook()
ws = wb.active; ws.title = "1. Metrics grid"
ws.sheet_properties.tabColor = "085041"
ws.merge_cells("A1:M1")
ws.cell(1, 1, "THE GRID — every metric × use case × data condition. Current formula cited from the code; "
              "the NEW formula is plain text in COLUMN K, and nothing reaches an engine that is not "
              "written here. Red K-cells are refusals. Drafted and verified by the workflow agents; "
              "assembled mechanically — no cell authored by Claude Code. PREVIEW until your verdicts."
).font = Font(bold=True, size=11, color="085041")
ws.cell(1, 1).alignment = WRAP; ws.row_dimensions[1].height = 50
COLS = [("Section", 9), ("Metric", 20), ("Use case", 26), ("Condition", 14), ("When this happens", 30),
        ("CURRENT formula (today, cited)", 40), ("Variables / inputs", 32), ("Source of each variable", 36),
        ("Availability", 18), ("Trap / how it goes wrong", 34),
        ("K — THE NEW FORMULA (plain text)", 64), ("YOUR VERDICT", 13), ("YOUR COMMENT", 28)]
for c, (t, w) in enumerate(COLS, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    cell = ws.cell(2, c, t); cell.fill, cell.font, cell.alignment, cell.border = HDR, WHITE, WRAP, THIN
assert openpyxl.utils.get_column_letter(11) == "K", "the new formula MUST land in column K"
ws.freeze_panes = "C3"

REFUSAL = re.compile(r"refus|not shown|no figure|absent|suppress|we do not|don't show|do not show|nothing is shown", re.I)
r_ = 3; last_sec = last_met = last_purp = None
for sec, metric, purpose, cond, case, rule in rows:
    if sec != last_sec:
        ws.merge_cells(start_row=r_, start_column=1, end_row=r_, end_column=13)
        cell = ws.cell(r_, 1, f"{sec} · {SECTION_TITLES.get(sec, sec)}")
        cell.fill, cell.font, cell.border = SECF, Font(bold=True, size=10, color="1B3A5C"), THIN
        last_sec, last_met = sec, None; r_ += 1
    if metric != last_met:
        ws.merge_cells(start_row=r_, start_column=1, end_row=r_, end_column=13)
        cell = ws.cell(r_, 1, metric)
        cell.fill, cell.font, cell.border = METF, Font(bold=True, size=10), THIN
        last_met, last_purp = metric, None; r_ += 1
    first_of_purpose = purpose != last_purp; last_purp = purpose
    sp = SPINE.get(metric.strip().lower(), {})
    vals = [sec, metric, purpose, cond, case,
            current_of(metric) if first_of_purpose else "", 
            vars_of(metric) if first_of_purpose else "",
            src_of(metric) if first_of_purpose else "",
            (sp.get("avail") or "") if first_of_purpose else "",
            (sp.get("trap") or "") if first_of_purpose else "",
            rule, "", ""]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(r_, c, v); cell.alignment, cell.border = WRAP, THIN
        if c in (6, 7, 8, 9, 10): cell.fill = GEN
        if c == 11: cell.fill = REF if REFUSAL.search(str(v)) else KFILL
        if c in (12, 13): cell.fill = YOURS
    r_ += 1

# ── tab 2: variables ← institutions (SnapTrade · Yodlee · Plaid, verbatim from the research) ─────
fc = spine["2. Field Crosswalk"]
hf = {fc.cell(4, c).value: c for c in range(1, fc.max_column + 1)}
v = wb.create_sheet("2. Variables from institutions")
v.sheet_properties.tabColor = "1B4D8F"
v.merge_cells("A1:G1")
v.cell(1, 1, "EVERY DATA ELEMENT — is it provided directly by SnapTrade, Yodlee or Plaid, and under what "
             "names. Copied VERBATIM from the field-crosswalk research (published specs + live-verified "
             "payloads); 'NOT PRESENT' means the standard does not carry it. A variable not provided by "
             "any of the three is derived or typed — its row says which."
).font = Font(bold=True, size=11, color="1B4D8F")
v.cell(1, 1).alignment = WRAP; v.row_dimensions[1].height = 46
VC = [("Variable / data element", 30), ("Provided directly by an institution?", 22),
      ("SnapTrade — potential names", 38), ("Yodlee — potential names", 34),
      ("Plaid — potential names", 40), ("Divergence risk", 36), ("YOUR COMMENT", 26)]
for c, (t, w) in enumerate(VC, start=1):
    v.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    cell = v.cell(2, c, t); cell.fill, cell.font, cell.alignment, cell.border = HDR, WHITE, WRAP, THIN
vr = 3
NOTP = re.compile(r"NOT PRESENT|—|derived", re.I)
for r in range(5, fc.max_row + 1):
    el = fc.cell(r, 1).value
    if not el: continue
    snap = str(fc.cell(r, hf.get("SnapTrade (brokerage-direct)", 9)).value or "")
    yod = str(fc.cell(r, hf.get("Yodlee v1.1", 5)).value or "")
    pl = " / ".join(x for x in [str(fc.cell(r, hf.get("Plaid Core Exchange (provider side, FDX subset)", 3)).value or ""),
                                 str(fc.cell(r, hf.get("Plaid Investments (recipient side)", 4)).value or "")] if x and x != "None")
    if not (snap or yod or pl):        # a section header row
        v.merge_cells(start_row=vr, start_column=1, end_row=vr, end_column=7)
        cell = v.cell(vr, 1, str(el)); cell.fill, cell.font, cell.border = SECF, Font(bold=True, size=10, color="1B3A5C"), THIN
        vr += 1; continue
    any_direct = not all(NOTP.search(x or "NOT PRESENT") for x in (snap, yod, pl))
    vals = [el, "Yes" if any_direct else "No — derived or typed", snap, yod, pl,
            fc.cell(r, hf.get("Divergence risk", 10)).value or "", ""]
    for c, val in enumerate(vals, start=1):
        cell = v.cell(vr, c, val); cell.alignment, cell.border = WRAP, THIN
        if c == 2: cell.fill = (PatternFill("solid", fgColor="EAF3EC") if any_direct else PatternFill("solid", fgColor="F6EEDC"))
        if c == 7: cell.fill = YOURS
    vr += 1
v.freeze_panes = "B3"

# ── open questions tab ──────────────────────────────────────────────────────────────────────────
oq = os.path.join(SRC, "OPEN-QUESTIONS.md")
q = wb.create_sheet("3. Open questions")
q.sheet_properties.tabColor = "A32D2D"
for c, (t, w) in enumerate([("#", 5), ("Question", 95), ("Your answer", 40)], start=1):
    q.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    cell = q.cell(2, c, t); cell.fill, cell.font, cell.border = HDR, WHITE, THIN
q.cell(1, 1, "Everything the agents could not decide as the CFA — your call.").font = Font(bold=True, size=11, color="A32D2D")
qi = 0
if os.path.exists(oq):
    for l in open(oq, encoding="utf-8"):
        m = re.match(r"\s*\d+[.)]\s*(.*)", l)
        if m:
            qi += 1; q.cell(qi + 2, 1, qi); q.cell(qi + 2, 2, m.group(1)); q.cell(qi + 2, 2).alignment = WRAP
            q.cell(qi + 2, 3).fill = YOURS

wb.save(OUT)
mets = len({(s, m) for s, m, *_ in rows})
print(f"{OUT}")
print(f"  grid: {mets} metrics · {len({(s,m,p) for s,m,p,*_ in rows})} use cases · {len(rows)} condition rows · K verified as column K")
print(f"  variables tab: {vr-3} rows · open questions: {qi}")
