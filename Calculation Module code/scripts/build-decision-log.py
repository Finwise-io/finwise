#!/usr/bin/env python3
"""PRINT THE DECISION LOG — one file, every decision he has made, wherever he made it.

Founder, 2026-09-28: *"one file in decision log folder as master source for all decisions"* and
*"have similar tabs <one for each parent screen + multi-screen> as master calculation file in decision
file to make it easy to read, and track."*

ALL DECISIONS is the authoritative list. The per-screen tabs are VIEWS of the same rows, printed in the
same pass, so they cannot drift from it — the failure this whole file exists to stop is one decision
recorded in two places that disagree. Nothing is edited in the workbook: it is printed from
moneykeel_decisions.json, the same rule as the calculation master and the working plan.
"""
import json, os, sys, datetime, importlib.util as _u
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
os.chdir(ROOT)
OUT = 'MoneyKeel/documents/decision-log/moneykeel_log_decisions_v01.xlsx'
# never write over a workbook he has open (2026-09-30) — it leaves his screen showing the OLD copy
from moneykeel_open_workbook import refuse_if_open
refuse_if_open(OUT)
TODAY = datetime.date.today().isoformat()

def _load(n):
    s = _u.spec_from_file_location(n, os.path.join(HERE, n + '.py'))
    m = _u.module_from_spec(s); s.loader.exec_module(m); return m
S = _load('moneykeel_decision_schema')
sys.path.insert(0, HERE)
from xlsx_readable import readable

doc = json.load(open(os.path.join(HERE, 'moneykeel_decisions.json')))
rows = doc['rows']

# ── ONE MAPPING, NOT TWO. Where a decision applies is authored in moneykeel_decisions_anchors.json;
#    each decision row also carried its own copy of `governs`, and nothing kept the two in step. On
#    2026-10-01 they had drifted: six anchors were corrected by hand after reading the code, and the
#    decision log went on printing the old, word-matched answer — DEC-0036 still pointing at the Net
#    worth screen, where the heads-up it governs has never been written. He has said what the decision
#    log is: *"Decision-log is one consolidated master."* A master that prints a second, staler copy of
#    a mapping is not one. So the row's `governs` is PRINTED from the anchors here, every build.
# ── HIS TWO CONTROLS ON WHAT MAY ENTER THE LOG (2026-10-01) ──────────────────────────────────────
#
# Control 1: a new decision must answer what is impacted, what changes, and the steps to apply it.
# Control 2: a new decision that repeats one he already made stops, and names the pair.
#
# Both bind NEW rows only and refuse the WRITE, not merely warn — his words: *"guard should return
# failed process."* What is already in the log predates them and is carried as a measured debt that
# prints every build. Full reasoning, and the two design calls, in moneykeel_decision_gate.py.
import moneykeel_decision_gate as GATE

_new = [r for r in rows if GATE.is_new(r)]
_refused = [(r, GATE.three_questions(r)) for r in _new]
_refused = [(r, m) for r, m in _refused if m]
if _refused:
    print('\nPROCESS FAILED — ' + GATE.REFUSAL)
    for _r, _m in _refused:
        print(f'\n   {_r["id"]}  {str(_r.get("his_words") or _r.get("plain") or "")[:90]}')
        for _q in _m:
            print(f'      unanswered: {_q}')
    print(f'\n   {len(_refused)} new decision(s) held back. Answer the three questions on the row '
          '— what it impacts, what changes, and the numbered steps to apply it — then build again.')
    sys.exit(1)

_pairs = [(n, p_, w, rr) for n, p_, w, rr in GATE.prior_decisions(rows) if GATE.is_new(n)]
_pairs = [(n, p_, w, rr) for n, p_, w, rr in _pairs
          if not GATE.resolved(n, p_)]
if _pairs:
    print('\nPROCESS FAILED — a new decision repeats one already recorded.')
    for _n, _p, _w, _rr in _pairs:
        print('\n   ' + GATE.supersede_question(_n, _p))
        print(f'      both govern: {" ".join(sorted(_rr)[:4])}')
        print(f'      same subject: {" ".join(sorted(_w)[:8])}')
    print('\n   Resolve each: if the new one replaces the old, mark the older row OLD, NOT IN USE and '
          'name the newer in its Replaces / Replaced by column; if they are genuinely different, record '
          'what is different on the new row. Then build again.')
    sys.exit(1)

_ap = os.path.join(HERE, 'moneykeel_decisions_anchors.json')
_anchors = json.load(open(_ap, encoding='utf8'))['rows']
_synced = 0
for _r in rows:
    _a = _anchors.get(_r.get('id'), {})
    if _a.get('governs') and _a['governs'] != _r.get('governs'):
        _r['governs'] = _a['governs']
        _synced += 1
    if _a.get('verified_by_reading'):
        _r['anchor_read_on'] = _a['verified_by_reading']
if _synced:
    json.dump(doc, open(os.path.join(HERE, 'moneykeel_decisions.json'), 'w', encoding='utf8'),
              indent=1, ensure_ascii=False)
    print(f'   governs synced from the anchors on {_synced} row(s) \u2014 one mapping, not two')

H = Font(bold=True, color='FFFFFF')
HFILL = PatternFill('solid', fgColor='14532D')       # the dark green, edge to edge

wb = openpyxl.Workbook(); wb.remove(wb.active)

# ── Read me ──────────────────────────────────────────────────────────────────────────────────────
ws = wb.create_sheet('Read me')
ws.append(['What this file is', 'Read this first'])
about = [
 ('THE ONE PLACE EVERY DECISION IS RECORDED',
  'Founder, 2026-09-28: "Rulings in master file needs to be one place to record all decisions, no matter '
  'where they are made. All the decisions sitting in change log, standards, requirements, Audit, and mockup '
  'files - record in separate file called decisions... After mistakes over and over, I no longer trust you." '
  'This is that file. The rulings tab was removed from the calculation master the same day, so there is one '
  'record and not two.'),
 ('WHY IT EXISTS — three misses, each one earned a column',
  '(1) His 2026-09-12 correction of income-end-to-end v07 lived in the design Changelog and the mock, in no '
  'ledger and no test — so the same defect was still live one screen over sixteen days later. (2) The design '
  'Changelog itself lost ten rows between its 2026-08-12 and 2026-08-14 versions, three of them his '
  'decisions, silently. (3) The rulings ledger this replaces began in August 2026 and held 43 of the 238 '
  'decisions his own Changelog records — so "check the rulings ledger" was advice that could not be relied on.'),
 ('HOW TO READ IT',
  'ALL DECISIONS is every row. The per-screen tabs are views of those same rows — one tab per parent screen, '
  'plus multiple_screens for a decision governing more than one, the same way the calculation master files a '
  'figure shown on three screens. process-and-standards holds decisions about how we work and about the '
  'standards themselves, which govern no single screen.'),
 ('NOTHING IS EDITED HERE',
  'This workbook is PRINTED from Calculation Module code/scripts/moneykeel_decisions.json by '
  'build-decision-log.py. A hand edit dies on the next build — the same rule as the calculation master and '
  'the working plan. To change a decision, change the source and rebuild.'),
 ('THE CONFIDENCE COLUMN — read it',
  'HIGH means his own words are on the row, quoted from the source. LOW means the source recorded what I '
  'BUILT and not what he SAID, so the row is my reconstruction and needs his wording or his confirmation. '
  'Across hundreds of rows some are inferred, and he must be able to see WHICH rather than trusting them all '
  'equally. Sort ALL DECISIONS by that column to find every row still owed.'),
 ('A ROW IS NEVER DELETED',
  'A decision that is reversed or replaced is SUPERSEDED and says by what. Deleting it would lose the fact '
  'that it was once true, which is what makes an old build readable.'),
 ('', ''),
 ('THE COLUMNS, and the failure each one answers', ''),
]
for a, b in about:
    ws.append([a, b])
for key, head, why in S.COLUMNS:
    ws.append([head, why])
ws.column_dimensions['A'].width = 46
ws.column_dimensions['B'].width = 118
for r in range(2, ws.max_row + 1):
    ws.cell(r, 1).font = Font(bold=True)
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')

def order(r):
    return (str(r.get('decided') or '9999'), str(r.get('id')))

def sheet(name, these, note=''):
    w = wb.create_sheet(name)
    if note:
        w.append([note]); w.append([])
    w.append(S.HEADINGS)
    hrow = w.max_row
    for c in range(1, len(S.HEADINGS) + 1):
        w.cell(hrow, c).font = H; w.cell(hrow, c).fill = HFILL
        w.cell(hrow, c).alignment = Alignment(wrap_text=True, vertical='top')
    for r in sorted(these, key=order):
        w.append([str(r.get(k) or '') for k in S.KEYS])
    w.freeze_panes = w.cell(hrow + 1, 2).coordinate
    return w

sheet('ALL DECISIONS', rows,
      'EVERY decision, oldest first. This is the authoritative list; the per-screen tabs below are views of these same rows.')
for tab in S.PARENT_TABS + ['process-and-standards']:
    these = [r for r in rows if r.get('tab') == tab]
    sheet(tab, these, S.TAB_ABOUT.get(tab, '') + f'   ·   {len(these)} decision(s)')

# every tab through the one formatter — measured widths, real row heights, frozen header, a filter
for w in wb.worksheets:
    if w.title == 'Read me':
        readable(w, widths={'A': 46, 'B': 118}, freeze='A2', filt=False, max_lines=14)
    else:
        readable(w, freeze=w.freeze_panes, filt=True)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('wrote', OUT)
print(f'   {len(rows)} decisions · {len(wb.worksheets)} tabs')
for tab in S.PARENT_TABS + ['process-and-standards']:
    print(f'   {tab:24s} {len([r for r in rows if r.get("tab") == tab]):4d}')
low = [r for r in rows if str(r.get('confidence', '')).startswith('LOW')]
print(f'   rows needing his wording or confirmation (confidence LOW): {len(low)}')
